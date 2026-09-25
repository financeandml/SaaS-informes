"""Secciones D–I y evidencia visual: el libro del DCF, los retratos, el Item 1A, el historial y los recortes de texto.

Cada prueba dice qué defecto reintroducido la hace fallar. Las que usan los adjuntos reales
se saltan si faltan; el libro del DCF de Netflix se busca en la carpeta del expediente.
"""

import json
import tempfile
import unittest
from datetime import date
from pathlib import Path

from tests import rutas_nflx
from tesis.heredado import dcf, historial, riesgos
from tesis.datos import expediente, gobierno, recortes
from tesis.datos.hechos import Capa, Contraste, Estado, Hecho, Periodo

RUTAS = rutas_nflx()
CARPETA = Path(json.loads((Path(__file__).parent / "expediente_nflx.json").read_text(encoding="utf-8"))["carpeta"])
DCF_NFLX = CARPETA / "Modelo_DCF_NFLX_Warrants_Co_2026-09-15.xlsx"


def _adjunto(nombre: str):
    return expediente.cargar_adjunto(next(r for r in RUTAS if r.name == nombre))


@unittest.skipUnless(RUTAS and DCF_NFLX.exists(), "sin expediente de NFLX o sin libro del DCF")
class LibroDCF(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.m = dcf.cargar(DCF_NFLX)

    def test_lee_los_bloques_por_rotulo(self):
        """Falla si el lector deja de encontrar por rótulo los escenarios, la sensibilidad, el reverse o los múltiplos."""
        m = self.m
        self.assertEqual([e.nombre for e in m.escenarios], ["pesimista", "base", "optimista"])
        base = m.escenario("base")
        self.assertAlmostEqual(base.wacc, 0.095)
        self.assertAlmostEqual(base.g, 0.0325)
        self.assertAlmostEqual(base.peso, 0.5)          # y no el «peso del valor terminal», que también empieza por «peso»
        self.assertEqual(base.anios[0], "2026E")
        self.assertIsNotNone(base.fila(r"FCFF"))
        self.assertAlmostEqual(base.dato(r"POR ACCIÓN — HOY")[1], 75.54, places=2)
        self.assertEqual(len(m.sensibilidad["ejes_g"]), 5)
        self.assertEqual(len(m.multiplos), 4)
        self.assertEqual(len(m.referencias_precio), 5)  # «Precio objetivo del inversor 100» es un dato, no una cabecera de bloque
        self.assertEqual(len(m.anclajes_objetivo), 4)
        # el libro del analista (18/09/2026) llegó guardado sin recalcular: el lector recalcula una copia con Excel y lo declara
        self.assertEqual([k for k in m.faltan if k != "valores"], [])
        if m.faltan.get("valores"):
            self.assertIsNotNone(m.recalculado, m.faltan["valores"])
            self.assertIn("copia recalculada", m.faltan["valores"])

    def test_un_libro_sin_valores_y_sin_excel_lo_declara(self):
        """Falla si un libro guardado sin recalcular, en una máquina sin Excel, deja pasar N/A en silencio o inventa valores."""
        import openpyxl
        from unittest import mock
        sin = dcf._sin_valores(openpyxl.load_workbook(str(DCF_NFLX), data_only=True), openpyxl.load_workbook(str(DCF_NFLX), data_only=False))
        if not sin:
            self.skipTest("este libro trae los valores guardados")
        # sin Excel y sin una copia recalculada de una emisión anterior (que ahora se reutiliza en vez de reescribirse)
        with mock.patch.object(dcf, "recalcular_con_excel", lambda ruta, copia: None), mock.patch.object(dcf, "_copia_vigente", lambda formulas, copia: None):
            m = dcf.cargar(DCF_NFLX)
        self.assertIn("valores", m.faltan)
        self.assertIsNone(m.recalculado)
        self.assertIn("no hay Excel", m.faltan["valores"])
        self.assertEqual(m.tabla_escenarios, [])

    def test_distingue_formula_de_valor_tecleado(self):
        """Falla si el lector deja de leer el libro dos veces (valores y fórmulas): el precio de mercado es tecleado y la capitalización es fórmula."""
        self.assertIsNone(self.m.supuesto("precio de mercado").celda.formula)
        self.assertTrue(self.m.supuesto("capitalizaci").celda.formula.startswith("="))

    def test_cuadre_caza_la_deuda_distinta_del_10q(self):
        """Falla si el cuadre da por buena una deuda del libro (16.655) que no es la deuda bruta contrastada (14.309) o deja de casar la caja."""
        fin = date(2026, 6, 30)
        i = Periodo.instante(fin)
        fy = Periodo.anual(date(2025, 12, 31))
        hechos = {("deuda_bruta", i): Hecho("deuda_bruta", i, 14_309_306_000.0, Estado.VALOR, Capa.SEC),
                  ("caja", i): Hecho("caja", i, 9_099_232_000.0, Estado.VALOR, Capa.SEC),
                  ("inversiones_cp", i): Hecho("inversiones_cp", i, 28_678_000.0, Estado.VALOR, Capa.SEC),
                  ("ingresos", fy): Hecho("ingresos", fy, 45_183_036_000.0, Estado.VALOR, Capa.SEC)}
        por_campo = {c.campo: c for c in dcf.cuadrar(self.m, hechos)}
        self.assertIs(por_campo["deuda_bruta"].contraste, Contraste.DISCREPANTE)
        self.assertIs(por_campo["caja_total"].contraste, Contraste.CONFIRMADO)
        self.assertIs(por_campo["ingresos"].contraste, Contraste.CONFIRMADO)
        self.assertIn("16.655", por_campo["deuda_bruta"].nota)

    def test_libro_sin_hojas_declara_lo_que_falta(self):
        """Falla si un libro sin hoja de escenarios pasa en silencio en vez de rellenar `faltan`."""
        import openpyxl
        with tempfile.TemporaryDirectory() as d:
            ruta = Path(d) / "vacio.xlsx"
            wb = openpyxl.Workbook(); wb.active.title = "Hoja1"; wb.active["A1"] = "nada"; wb.save(ruta)
            m = dcf.cargar(ruta)
        for clave in ("supuestos", "escenario_base", "sensibilidad", "reverse", "resumen"):
            self.assertIn(clave, m.faltan)


@unittest.skipUnless(RUTAS, "sin expediente de NFLX")
class RetratosYFichas(unittest.TestCase):
    def test_cinco_ejecutivos_con_retrato_y_trayectoria(self):
        """Falla si el retrato deja de asociarse a su ficha por posición o si las viñetas de trayectoria no se leen."""
        proxy = _adjunto("Netflix-2026-Proxy-Statement.pdf")
        with tempfile.TemporaryDirectory() as d:
            g = gobierno.construir(expediente.Expediente("NFLX", [proxy]), Path(d))
            self.assertEqual(len(g.ejecutivos), 5)
            self.assertTrue(all(e.foto is not None and e.foto.ruta.exists() for e in g.ejecutivos))
            neumann = next(e for e in g.ejecutivos if e.nombre == "Spencer Neumann")
            self.assertEqual(neumann.foto.pagina, 40)
            self.assertIn("CFO of Netflix (since 2019)", neumann.trayectoria)
            self.assertEqual(len(g.consejeros), 13)
            peters = next(c for c in g.consejeros if c.nombre == "Greg Peters")
            self.assertTrue(peters.cargo.startswith("CO-CHIEF EXECUTIVE OFFICER"))   # el cargo de tres líneas, entero
            hastings = next(c for c in g.consejeros if c.nombre == "Reed Hastings")
            self.assertIn("no se presenta a la reelección", hastings.nota)


@unittest.skipUnless(RUTAS, "sin expediente de NFLX")
class Item1A(unittest.TestCase):
    def test_epigrafes_en_negrita_con_familia_y_superados(self):
        """Falla si el lector pierde epígrafes (texto sin espacios, negrita mal medida) o deja de rotular como superado el riesgo WBD."""
        exp = expediente.Expediente("NFLX", [_adjunto("99482238-46b2-4d0d-b292-40e6781bdf03.pdf"), _adjunto("65ef36cc-6598-4deb-a42b-8c59a8e31fd3.pdf")])
        r = riesgos.construir(exp)
        self.assertEqual(r.paginas, (7, 18))
        self.assertGreaterEqual(len(r.riesgos), 30)
        self.assertLessEqual(len(r.riesgos), 45)          # si la negrita se mide mal, cada frase de cada párrafo pasa por epígrafe
        epigrafes = [x.epigrafe for x in r.riesgos]
        self.assertIn("Our stock price is volatile.", epigrafes)
        self.assertIn("If our efforts to attract and retain members are not successful, our business will be adversely affected.", epigrafes)
        self.assertTrue(all(" " in x.epigrafe for x in r.riesgos))     # espacios reconstruidos: «If we» y no «Ifw e»
        self.assertTrue(all(x.familia in riesgos.FAMILIAS and x.motivo for x in r.riesgos))
        wbd = [x for x in r.riesgos if "WBD" in x.epigrafe]
        self.assertTrue(wbd and all("superado" in x.nota for x in wbd))


@unittest.skipUnless(RUTAS, "sin expediente de NFLX")
class HistorialYRecortes(unittest.TestCase):
    def test_consenso_frente_a_real_contrastado(self):
        """Falla si la tabla de la portada de la transcripción deja de leerse o si el «real» no se contrasta con el hecho del informe."""
        exp = expediente.Expediente("NFLX", [_adjunto("Netflix-Inc-_Earnings-Call_2026-07-16T00_00_00_English-1.pdf"), _adjunto("FINAL-Q2-26-Shareholder-Letter.pdf")])
        p2 = Periodo.de_meses(date(2026, 6, 30), 3)
        hechos = {("ingresos", p2): Hecho("ingresos", p2, 12_559_938_000.0, Estado.VALOR, Capa.SEC),
                  ("bpa_diluido", p2): Hecho("bpa_diluido", p2, 0.80, Estado.VALOR, Capa.SEC, unidad="USD/acción")}
        h = historial.construir(exp, hechos)
        self.assertEqual(h.proveedor, "S&P Global Market Intelligence")
        self.assertEqual([s.periodo for s in h.sorpresas], ["3T25", "4T25", "1T26", "2T26", "2T26"])
        s = next(s for s in h.sorpresas if s.periodo == "3T25")
        self.assertAlmostEqual(s.sorpresa, -0.1571)
        self.assertEqual({s.contraste for s in h.sorpresas if s.periodo == "2T26"}, {Contraste.CONFIRMADO})
        self.assertEqual(len(h.frases_guia), 3)
        self.assertTrue(all("$12.6B" in c.texto or "forecast" in c.texto for c in h.frases_guia))   # «$12.6B» no parte la frase

    def test_recorte_de_lineas_exige_el_ancla(self):
        """Falla si se produce un recorte de una página que no contiene el ancla, o si el recorte no incluye la línea del ancla."""
        carta = _adjunto("FINAL-Q2-26-Shareholder-Letter.pdf")
        with tempfile.TemporaryDirectory() as d:
            r = recortes.recortar_lineas(carta, 2, ["we are narrowing our revenue forecast to $51.0-$51.4B"], Path(d), ["prueba"], "p")
            self.assertIsNotNone(r)
            self.assertTrue(r.ruta.exists() and r.rectangulo[3] - r.rectangulo[1] < 200)   # unas pocas líneas, no la página entera
            self.assertIsNone(recortes.recortar_lineas(carta, 2, ["texto que no está en la carta"], Path(d), ["prueba"], "q"))


if __name__ == "__main__":
    unittest.main()
