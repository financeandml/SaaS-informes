"""Los 78 fallos de la auditoría del 27/09/2026, uno por prueba (docs/fases/A.md).

Cada prueba afirma el estado CORRECTO y empieza en `xfail` estricto: mientras el fallo esté, falla y cuenta como
esperado; cuando una parte lo arregla, la prueba pasa, `strict` lo convierte en error y hay que quitar su marca. Así el
contador `pytest -m auditoria` dice cuántos de los 78 quedan, y ninguno se da por arreglado sin su prueba.

Las cuatro emisiones (AAPL, NFLX, ORCL y QCOM a 27/09/2026) se hacen una vez, sin red, con lo congelado en
`tests/fixtures/auditoria/` (`tests/grabar_auditoria.py`) y las entradas de los analistas de prueba de ese día. Los
adjuntos del analista no se versionan: sin ellos (`WC_DATOS/adjuntos/<TICKER>`) estas pruebas se saltan.
"""

from __future__ import annotations

import json
import os
import re
import shutil
import sys
import tempfile
import unittest
import unittest.mock
from datetime import date
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent))
from informe_html import Informe                                        # noqa: E402

RAIZ = Path(__file__).resolve().parents[1]
F = RAIZ / "tests" / "fixtures" / "auditoria"
TICKERS = ("AAPL", "NFLX", "ORCL", "QCOM")
DIA = "2026-09-27"

pytestmark = [pytest.mark.auditoria, pytest.mark.lenta]


def falla(numero: int, motivo: str):
    """El fallo [numero] de la auditoría, aún sin arreglar."""
    return pytest.mark.xfail(strict=True, reason=f"[{numero}] {motivo}")


def _adjuntos(t: str) -> Path:
    from tesis import entorno
    return entorno.carpeta("adjuntos") / t


HAY = F.exists() and all(_adjuntos(t).exists() for t in TICKERS)
_E: dict = {}


def setUpModule():
    if not HAY:
        return
    from tesis import entorno
    from tesis.fuentes import precio, sec
    _E["previo"] = (dict(os.environ), sec.CACHE, precio.CACHE_BOLSA, precio.SOLO_CACHE)
    _E["tmp"] = Path(tempfile.mkdtemp())
    os.environ.update(WC_SEC_CONTACTO="", WC_PRECIO_FUENTE="nasdaq")
    sec.CACHE, precio.CACHE_BOLSA, precio.SOLO_CACHE = F / "cache_sec", F / "cache_bolsa", True
    sys.path.insert(0, str(RAIZ))
    import emitir
    _E["codigo"], _E["inf"], _E["aud"], _E["base"] = {}, {}, {}, {}
    for t in TICKERS:
        salida = _E["tmp"] / t
        _E["codigo"][t] = emitir.main([t, "--carpeta", str(_adjuntos(t)), "--entradas", str(F / t / "entradas.json"),
                                       "--fecha", DIA, "--salida", str(salida)])
        base = salida / f"{t}_tesis_{DIA}"
        _E["base"][t] = base
        _E["inf"][t] = Informe(base.with_suffix(".html").read_text(encoding="utf-8"))
        _E["aud"][t] = json.loads(base.with_suffix(".auditoria.json").read_text(encoding="utf-8"))
    del entorno


def tearDownModule():
    if "previo" not in _E:
        return
    from tesis.fuentes import precio, sec
    env, sec.CACHE, precio.CACHE_BOLSA, precio.SOLO_CACHE = _E["previo"]
    os.environ.clear()
    os.environ.update(env)
    shutil.rmtree(_E["tmp"], ignore_errors=True)


def inf(t: str) -> Informe:
    return _E["inf"][t]


def aud(t: str) -> dict:
    return _E["aud"][t]


def bloqueos(t: str):
    return aud(t).get("bloqueos") or []


def avisos(t: str):
    return aud(t).get("avisos") or []


def num(texto: str) -> float:
    """«1.234,5 %» → 1234.5 (es-ES, signo menos tipográfico)."""
    t = texto.replace("−", "-").replace(".", "").replace(",", ".")
    return float(re.search(r"-?\d+(?:\.\d+)?", t).group(0))


def ultima(t: str, cuadro: str, fila: str) -> str:
    return inf(t).fila(cuadro, fila)[-1]


def definido(xlsx: Path, nombre: str) -> float:
    from openpyxl import load_workbook
    wb = load_workbook(xlsx, data_only=False)
    hoja, celda = list(wb.defined_names[nombre].destinations)[0]
    return float(wb[hoja][celda.replace("$", "")].value)


@unittest.skipUnless(HAY, "sin fixtures de la auditoría o sin los adjuntos del analista")
class Mercado(unittest.TestCase):
    def test_01_sesion_del_precio(self):
        for t in TICKERS:
            self.assertFalse(inf(t).celda("Ficha", "Cierre anterior", "Valor").startswith("N/A"), t)
            self.assertIn("25/09/2026", inf(t).celda("Ficha", "Volumen", "Valor"), t)

    def test_02_historico_de_evidencia_hasta_la_valoracion(self):
        for t in TICKERS:
            ruta = _E["base"][t].parent / f"{t}_tesis_{DIA}_recortes" / "api_nasdaq_historico.json"
            filas = json.loads(ruta.read_text(encoding="utf-8"))["data"]["tradesTable"]["rows"]
            self.assertEqual(filas[0]["date"], "09/25/2026", t)

    def test_42_ultima_operacion_de_la_cadena(self):
        self.assertIn("última operación: 137,10 USD el 25/09/2026", inf("ORCL").texto)

    def test_66_fecha_de_nasdaq_no_manda(self):
        self.assertIn("última operación: 201,97 USD el 25/09/2026", inf("QCOM").texto)
        self.assertNotIn("es del 24/09/2026, no del día de emisión", json.dumps(aud("QCOM"), ensure_ascii=False))


@unittest.skipUnless(HAY, "sin fixtures de la auditoría o sin los adjuntos del analista")
class Extraccion(unittest.TestCase):
    @falla(3, "«In millions, except shares in thousands» se aplica a toda la tabla")
    def test_03_escala_por_fila(self):
        for b in bloqueos("AAPL"):
            m = re.search(r"SEC ([\d.,]+) frente a .*? = ([\d.,]+)", b)
            if m and num(m.group(2)):
                self.assertNotAlmostEqual(num(m.group(1)) / num(m.group(2)), 1000, delta=1, msg=b)

    @falla(4, "columna del año anterior e intangibles de distinto alcance emparejados")
    def test_04_emparejamientos(self):
        texto = " ".join(bloqueos("AAPL"))
        for malo in ("8.268.000.000", "84.118", "Intangibles"):
            self.assertNotIn(malo, texto)

    @falla(5, "el split detectado sigue bloqueando")
    def test_05_split_no_bloquea(self):
        self.assertFalse([b for b in bloqueos("NFLX") if b.startswith("Discrepancia abierta")])

    @falla(49, "guía anterior al split sin reexpresar")
    def test_49_guia_reexpresada(self):
        fila = inf("NFLX").fila("Guía de la compañía", "3T FY24 · BPA diluido")
        self.assertEqual(fila[:2], ["0,51", "0,54"])


@unittest.skipUnless(HAY, "sin fixtures de la auditoría o sin los adjuntos del analista")
class TresEstados(unittest.TestCase):
    @falla(6, "un «no disponible» de la fuente se imprime como «no paga dividendo»")
    def test_06_no_disponible_no_es_cero(self):
        self.assertNotIn("no paga dividendo", inf("ORCL").fila("Fechas clave", "Dividendo")[0])

    @falla(48, "dividendos de NFLX: N/A donde hay un cero declarado; fuente Nasdaq para una política del 10-K")
    def test_48_dividendos_nflx(self):
        fila = inf("NFLX").fila("Flujo de caja", "Dividendos pagados")
        self.assertNotIn("N/A", fila)
        self.assertIn("10-K", inf("NFLX").fila("Fechas clave", "Dividendo")[1])

    @falla(52, "ORCL valorada sin dividendo (DPA 0)")
    def test_52_dpa_orcl(self):
        self.assertAlmostEqual(definido(_E["base"]["ORCL"].with_suffix(".motor.xlsx"), "dpa_horizonte"), 2.00, delta=0.01)

    @falla(24, "DPA del horizonte = suma de los cuatro últimos pagos, no el vigente")
    def test_24_dpa_vigente(self):
        self.assertAlmostEqual(definido(_E["base"]["QCOM"].with_suffix(".motor.xlsx"), "dpa_horizonte"), 3.68, delta=0.01)

    @falla(59, "amortización de deuda N/A con conceptos XBRL publicados")
    def test_59_amortizacion_de_deuda_qcom(self):
        fila = inf("QCOM").fila("Flujo de caja", "Amortización de deuda")
        cab = inf("QCOM").cabecera("Flujo de caja")
        self.assertEqual(fila[cab.index("2023") - 1], "−1.446")

    @falla(56, "interés en corto de NYSE sin fuente (FINRA)")
    def test_56_corto_orcl(self):
        self.assertNotIn("sin dato", " ".join(inf("ORCL").fila("Lista de comprobación", "Interés en corto")))

    @falla(67, "«pendiente de la API» cuando la fuente dice que no cubre el valor")
    def test_67_motivo_de_la_fuente(self):
        for t in TICKERS:
            self.assertNotIn("pendiente de la API", inf(t).texto, t)


@unittest.skipUnless(HAY, "sin fixtures de la auditoría o sin los adjuntos del analista")
class DeudaYAcciones(unittest.TestCase):
    def test_07_papel_comercial(self):
        i = inf("AAPL")
        self.assertEqual(i.fila("Balance", "Deuda bruta")[i.cabecera("Balance").index("2025") - 1], "98.657")

    def test_08_una_cifra_de_acciones(self):
        for t in TICKERS:
            i = inf(t)
            e = num(re.search(r"E = precio × acciones = ([\d.]+) M", i.fila("Construcción del WACC", "Peso de los fondos propios")[1]).group(1))
            precio = num(i.celda("Ficha", "Precio", "Valor"))
            acciones = num(i.fila("Puente", "Acciones diluidas")[0])
            self.assertAlmostEqual(e / precio, acciones, delta=1.0, msg=t)

    def test_45_deuda_neta_aapl(self):
        self.assertNotIn("la deuda neta era de 19.948", inf("AAPL").texto)
        self.assertEqual(inf("AAPL").fila("Puente", "(−) Deuda financiera")[0], "84.344")

    def test_54_checklist_deuda_vigente(self):
        self.assertIn("2,6x", " ".join(inf("ORCL").fila("Lista de comprobación", "Deuda neta / EBITDA")))

    def test_58_dilucion_adicional(self):
        acciones = num(inf("QCOM").fila("Puente", "Acciones diluidas")[0])
        self.assertTrue(acciones > 1069.5 or any("dilu" in a.lower() for a in avisos("QCOM")))

    def test_69_hecho_material(self):
        self.assertTrue(any("3.02" in a for a in avisos("QCOM")))


@unittest.skipUnless(HAY, "sin fixtures de la auditoría o sin los adjuntos del analista")
class Motor(unittest.TestCase):
    ROE = {"AAPL": "171,4 %", "NFLX": "42,8 %", "ORCL": "54,3 %", "QCOM": "23,3 %"}

    def test_13_roe(self):
        for t, v in self.ROE.items():
            self.assertEqual(ultima(t, "Rentabilidad y eficiencia", "ROE"), v, t)

    def test_43_roa(self):
        self.assertEqual(ultima("AAPL", "Rentabilidad y eficiencia", "ROA"), "30,9 %")
        self.assertEqual(ultima("ORCL", "Rentabilidad y eficiencia", "ROA"), "7,9 %")

    def test_14_roic_una_definicion(self):
        for t in TICKERS:
            titulos = inf(t)._titulos
            for rotulo, filas in titulos.items():
                if "Rentabilidad y eficiencia" in rotulo:
                    fila = next(f for f, c in zip(filas, inf(t).cuadros[rotulo]) if c and c[0] == "ROIC")
                    self.assertTrue(all("inversiones a corto plazo" in x.lower() for x in fila), t)

    def test_15_cagr_cinco_anios(self):
        for t, v in {"AAPL": "8,7 %", "QCOM": "13,5 %"}.items():
            self.assertEqual(inf(t).celda("DCF inverso", "CAGR de ingresos implícito", "Histórico 5 años"), v, t)

    def test_44_cagr_cinco_anios(self):
        for t, v in {"NFLX": "12,6 %", "ORCL": "10,7 %"}.items():
            self.assertEqual(inf(t).celda("DCF inverso", "CAGR de ingresos implícito", "Histórico 5 años"), v, t)

    def test_16_trimestre_de_caja(self):
        for t in TICKERS:
            balance = re.search(r"balance a (\d\d/\d\d/\d{4})", inf(t).fila("Puente", "(−) Deuda financiera")[1]).group(1)
            self.assertIn(balance, inf(t).fila("Supuestos generales", "Periodo parcial")[1], t)

    def test_17_sbc_no_se_resta(self):
        for t in TICKERS:
            i = inf(t)
            col = lambda r: num(i.fila("Escenario base: drivers", r)[1])                  # noqa: E731
            paquete = sum(num(f[2]) for f in i.cuadro("Escenario base: drivers") if f and f[0].startswith("(−) ") and f[0] not in
                          ("(−) Impuestos", "(−) Capex", "(−) Δ fondo de maniobra", "(−) SBC"))
            self.assertAlmostEqual(col("FCFF"), col("NOPAT") + col("(+) D&A") - col("(−) Capex") - col("(−) Δ fondo") - paquete,
                                   delta=2, msg=t)

    def test_75_spec_sbc(self):
        texto = (RAIZ / "docs" / "spec" / "05_motor_dcf.md").read_text(encoding="utf-8")
        seccion = texto[texto.index("## 3."):texto.index("## 4.")]
        self.assertIn("GAAP", seccion)

    def test_18_cierre_52_53(self):
        self.assertIn("26/09/2026", inf("AAPL").fila("Supuestos generales", "Periodo parcial")[1])

    def test_60_cierre_qcom(self):
        self.assertIn("27/09/2026", inf("QCOM").fila("Supuestos generales", "Periodo parcial")[1])

    def test_19_bpa_preferentes(self):
        self.assertEqual(inf("ORCL").celda("Estado de resultados", "BPA diluido", "4T FY26"), "1,45")

    def test_53_per_orcl(self):
        self.assertEqual(inf("ORCL").fila("Múltiplos sobre", "PER (TTM)")[0], "21,5x")

    def test_20_otros_ingresos(self):
        i, col = inf("ORCL"), "2026"
        v = lambda r: num(i.celda("Estado de resultados", r, col))              # noqa: E731
        # «EBIT (»: la fila se busca por prefijo y «EBIT» a secas casaba con la del EBITDA, que va antes
        self.assertAlmostEqual(v("EBIT (") + v("Gastos financieros") + v("Otros ingresos"), v("Resultado antes"), delta=1)

    def test_21_geografias(self):
        self.assertNotIn("geografías suman", inf("ORCL").texto)

    def test_22_ratios_no_significativos(self):
        self.assertNotIn("−24,8 %", inf("ORCL").fila("Flujo de caja", "Retribución / FCF"))

    def test_23_no_recurrentes(self):
        self.assertTrue(any("no recurrente" in a for a in avisos("NFLX")))


@unittest.skipUnless(HAY, "sin fixtures de la auditoría o sin los adjuntos del analista")
class Documentos(unittest.TestCase):
    @falla(10, "presentaciones y transcripciones no citables")
    def test_10_citables(self):
        self.assertFalse([b for b in bloqueos("QCOM") if "no disponible para verificar" in b])

    def test_25_consejo(self):
        for t in ("AAPL", "QCOM", "ORCL"):
            self.assertFalse([b for b in bloqueos(t) if b.startswith("Gobierno · consejo")], t)

    @falla(46, "retribución de AAPL del ejercicio anterior")
    def test_46_retribucion_ultimo_anio(self):
        fila = inf("AAPL").fila("Retribución de los ejecutivos", "Tim Cook")
        self.assertEqual((fila[0], fila[-1]), ("2025", "74.294.811"))

    @falla(50, "biografías mezcladas y cargos con restos de maquetación")
    def test_50_biografias(self):
        self.assertNotIn("Director Logo", inf("NFLX").html)
        sarandos = inf("NFLX").texto[inf("NFLX").texto.find("Ted Sarandos"):][:1500]
        self.assertNotIn("Microsoft", sarandos)

    @falla(26, "Item 1A de ORCL con un solo epígrafe")
    def test_26_item_1a(self):
        rotulo = next(k for k in inf("ORCL").cuadros if "Riesgos del Item 1A" in k)
        self.assertGreaterEqual(int(re.search(r"\((\d+) epígrafes", rotulo).group(1)), 20)

    @falla(57, "entradas válidas del analista descartadas (Item 1A de los pilares)")
    def test_57_riesgo_1a_aceptado(self):
        self.assertFalse([b for b in bloqueos("ORCL") if "riesgo_1a" in b])

    def test_27_fundacion(self):
        self.assertEqual(inf("ORCL").celda("Ficha", "Fundación", "Valor"), "1977")

    def test_28_guia_en_texto(self):
        self.assertFalse(inf("ORCL").cuadro("Objetivos vigentes")[1][0].startswith("N/A"))

    def test_70_na_bloqueado_aapl(self):
        self.assertFalse([b for b in bloqueos("AAPL") if "donde no puede haberlo" in b or "Texto técnico" in b])

    def test_71_na_bloqueado_orcl(self):
        self.assertFalse([b for b in bloqueos("ORCL") if "Apartado 2:" in b or "Apartado portada" in b])

    @falla(68, "los PDF con nombre UUID se muestran por su nombre de fichero")
    def test_68_rotulo_de_adjunto(self):
        from tesis.web import saas
        filas = saas.clasificar("ORCL")["adjuntos"]
        uuid = [f for f in filas if re.match(r"[0-9a-f]{8}-", f["fichero"])]
        self.assertTrue(uuid and all(f.get("rotulo") and not f["rotulo"].startswith(f["fichero"][:8]) for f in uuid))


class Esquema(unittest.TestCase):
    @falla(74, "sin campos para aportar a mano el gobierno que el parser no encuentra")
    def test_74_gobierno_manual(self):
        texto = (RAIZ / "docs" / "spec" / "04_entradas.yaml").read_text(encoding="utf-8")
        self.assertIn("gobierno.manual", texto)

    @falla(11, "el folio impreso «Apple Inc. | 2025 Form 10-K | 8» no se reconoce")
    def test_11_folio_de_pie_con_separadores(self):
        from tesis.datos import tablas_html
        # el folio impreso empieza en 3 (portada e índice sin él): si no se lee, sale el índice físico 1…8
        paginas = [f"texto de la página\nApple Inc. | 2025 Form 10-K | {k}" for k in range(3, 11)]
        self.assertEqual(tablas_html.folios_por_pagina(paginas), [str(k) for k in range(3, 11)])

    @falla(36, "una referencia a un grupo del escenario pasa el paso 6 y rompe la plantilla")
    def test_36_driver_de_grupo(self):
        from tesis.entradas import Entradas, comprobar_paso6
        e = Entradas({"esc": {"pesimista": {"paquete": {"desfase_contenido": {"1": 3}}}},
                      "bear": {"disparadores": [{"descripcion": "x", "metrica": "m", "umbral": 1, "unidad": "%", "plazo": "FY27",
                                                  "driver": "esc.pesimista.paquete"}]}})
        faltas = comprobar_paso6(e, {}, 0.9, lambda c: True)
        self.assertTrue(any("esc.pesimista.paquete" in f for f in faltas))

    @falla(35, "la fila SOM por defecto lleva texto técnico")
    def test_35_som_por_defecto(self):
        from tesis.entradas import Entradas
        from tesis.plantillas import parte_e
        from tesis.plantillas.informe import Cuadros
        literal = "Our platforms power devices across industries."
        e = Entradas({"mercado": {"cifras": [{"etiqueta": "Mercado", "clase": "TAM", "valor": 1000, "unidad": "millones USD", "anio": 2029,
                                              "metodo": "top_down", "evidencia": [{"doc": "10-K", "texto": literal, "texto_es": "Sus plataformas impulsan dispositivos."}]}]}})
        d = parte_e.ParteE()
        parte_e._mercado(d, Cuadros(), e, {"10-K": literal}, 0.9, None, 45e9, "2025", 2025)
        textos = " ".join(c.texto for f in d.cuadros["mercado"].filas for c in f.celdas)
        self.assertNotIn("sección C", textos)
        self.assertNotIn("companyfacts", textos)

    def test_72_erp(self):
        """[72] Falla si la ERP de referencia que delegó el analista (Damodaran a 01/09/2026) no se propone."""
        from tesis.entradas import proponer
        p = proponer._REGISTRO["wacc.erp"](proponer.Contexto("QCOM", date(2026, 9, 27), {}))
        self.assertEqual(getattr(p, "valor", None), 4.14)

    @falla(76, "la rúbrica da por hechas cosas que no lo están (SOTP, hoja «Reverse DCF»)")
    def test_76_rubrica(self):
        import importlib.util
        import yaml
        criterios = yaml.safe_load((RAIZ / "docs" / "spec" / "08_rubrica.yaml").read_text(encoding="utf-8"))["criterios"]
        sotp = [c for c in criterios if "SOTP" in c["criterio"] or "suma de partes" in c["criterio"].lower()]
        self.assertTrue(sotp)
        hecho = importlib.util.find_spec("tesis.motor.sotp") is not None
        self.assertTrue(all((c["estado"] == "si") == hecho for c in sotp))
        self.assertTrue(hecho)

    def test_77_f12_cerrada(self):
        self.assertTrue((RAIZ / "docs" / "fases" / "F12.md").exists())

    def test_78_raiz_vigilada(self):
        """[78] Falla si la batería deja de vigilar los ficheros de la raíz del repositorio."""
        import conftest
        self.assertIn("CLAUDE.md", conftest._foto_raiz())

    @falla(40, "SOTP sin implementar (decisión 5: se implementa)")
    def test_40_sotp(self):
        from tesis.motor import sotp                                                 # noqa: F401


@unittest.skipUnless(HAY, "sin fixtures de la auditoría o sin los adjuntos del analista")
class Puerta(unittest.TestCase):
    @falla(29, "bloqueos duplicados y recuentos distintos en checklist, Hoja 0 y log")
    def test_29_recuentos(self):
        for t in TICKERS:
            self.assertEqual(len(bloqueos(t)), len(set(bloqueos(t))), t)
        otros = len([b for b in bloqueos("ORCL") if not b.startswith("Discrepancia abierta")])
        self.assertIn(f"{otros} bloqueos", " ".join(inf("ORCL").fila("Lista de comprobación", "0 discrepancias")))

    @falla(61, "QCOM cuenta 4 bloqueos siendo 3 distintos")
    def test_61_recuento_qcom(self):
        self.assertIn("3 bloqueos", " ".join(inf("QCOM").fila("Lista de comprobación", "0 discrepancias")))

    @falla(30, "«Discrepancias: Ninguna» con una identidad que no cuadra")
    def test_30_identidades(self):
        for t in ("QCOM", "ORCL"):
            self.assertNotIn("Discrepancias\nNinguna.", inf(t).texto, t)

    @falla(31, "catalizadores «2026-T4» no cuentan como fechados")
    def test_31_catalizadores(self):
        self.assertFalse(" ".join(inf("ORCL").fila("Lista de comprobación", "Al menos un catalizador")).count("0 catalizadores"))
        self.assertIn("3 catalizadores", " ".join(inf("QCOM").fila("Lista de comprobación", "Al menos un catalizador")))

    @falla(32, "decisiones atribuidas al analista que no tomó")
    def test_32_atribuciones(self):
        self.assertNotIn("por decisión del analista (entradas, sotp.aplica", inf("ORCL").texto)

    @falla(33, "lo confirmado sigue rotulado «propuesta»")
    def test_33_confirmado(self):
        for t in TICKERS:
            self.assertNotIn("propuesta: la confirma el analista", inf(t).fila("Ficha", "Empleados")[1], t)

    @falla(34, "columna «Actual» de los KPI vacía teniendo el dato")
    def test_34_kpi_actual(self):
        for t in TICKERS:
            filas = inf(t).cuadro("Indicadores que se vigilan")[1:]
            self.assertTrue(any(f[1] not in ("—", "") for f in filas if len(f) > 1), t)

    @falla(37, "el párrafo factual habla de dividendos a quien no los paga")
    def test_37_parrafo_factual(self):
        self.assertNotIn("entre dividendos y recompras", inf("NFLX").texto)

    @falla(38, "RONIC sin publicar y la proyección completa no suma")
    def test_38_supuestos_visibles(self):
        for t in TICKERS:
            self.assertIn("RONIC", " ".join(" ".join(f) for f in inf(t).cuadro("Escenarios y valor razonable")), t)
        filas = [f[0].lower() for f in inf("NFLX").cuadro("Proyección completa") if f]
        self.assertTrue(any("contenido" in f for f in filas))

    @falla(39, "Excel sin hoja «Reverse DCF» con fórmulas ni tercer valor terminal")
    def test_39_excel(self):
        from openpyxl import load_workbook
        wb = load_workbook(_E["base"]["QCOM"].with_suffix(".motor.xlsx"))
        formulas = [c for fila in wb["Reverse DCF"].iter_rows() for c in fila if isinstance(c.value, str) and c.value.startswith("=")]
        self.assertTrue(formulas)

    @falla(41, "emitir.py devuelve 1 en borrador (decisión 4: 2)")
    def test_41_borrador_devuelve_2(self):
        for t in TICKERS:
            self.assertEqual(_E["codigo"][t], 2, t)


@unittest.skipUnless(HAY, "sin fixtures de la auditoría o sin los adjuntos del analista")
class Periodos(unittest.TestCase):
    @falla(12, "etiquetas de trimestre de calendario («3T25») junto a las fiscales")
    def test_12_trimestres_fiscales(self):
        for t in TICKERS:
            self.assertFalse(re.findall(r"\b[1-4]T\d{2}\b", inf(t).texto), t)

    @falla(47, "AAPL: marcas mal capitalizadas y RSU como «ejercicio de opciones»")
    def test_47_texto_aapl(self):
        self.assertNotIn("IPhone", inf("AAPL").texto)
        self.assertNotIn("Ejercicio de opciones", inf("AAPL").texto)


@unittest.skipUnless(HAY, "sin fixtures de la auditoría o sin los adjuntos del analista")
class MercadoSecundario(unittest.TestCase):
    @falla(9, "caché de EDGAR de días anteriores sin refrescar ni avisar")
    def test_09_cache_edgar(self):
        for t in ("QCOM", "ORCL"):
            self.assertTrue(any("caché" in a and "días" in a for a in avisos(t)), t)

    @falla(55, "13G/A al 0 % por reorganización leído como desinversión")
    def test_55_vanguard(self):
        self.assertNotIn("Ya no superan el 5 %", inf("ORCL").texto)

    @falla(51, "participación de la proxy fechada con la junta; BPA 2022 de NFLX mal redondeado")
    def test_51_fechas_de_participacion(self):
        self.assertEqual(inf("NFLX").celda("Estado de resultados", "BPA diluido", "2022"), "1,00")
        fila = next(f for f in inf("NFLX").cuadro("Accionistas con 5 %") if "BlackRock" in f[0])
        self.assertIn("2023", fila[-1])

    @falla(62, "el mismo hecho con dos valores (ingresos 4T FY25; consenso)")
    def test_62_un_hecho_un_valor(self):
        fila = inf("QCOM").fila("Guía de la compañía", "4T FY25 · Ingresos")
        self.assertEqual(num(fila[1]), num(inf("QCOM").celda("Estado de resultados", "Ingresos", "4T FY25")))


@unittest.skipUnless(HAY, "sin fixtures de la auditoría o sin los adjuntos del analista")
class AvisosDeEntradas(unittest.TestCase):
    """Tramo 5: avisan, no bloquean (decisión del 26/09/2026)."""

    def _avisos_asistente(self, t):
        from tesis.entradas import asistente
        datos = json.loads((F / t / "entradas.json").read_text(encoding="utf-8"))
        return asistente.validar(t, date(2026, 9, 27), datos)[1]

    @falla(63, "evidencias que no dicen lo que afirman y traducciones que quitan el modal")
    def test_63_evidencias(self):
        self.assertTrue(any("catalizador" in a.lower() and "evidencia" in a.lower() for a in self._avisos_asistente("QCOM")))

    @falla(64, "invalidaciones con listas distintas y TAM/SOM que no son mercado")
    def test_64_coherencia(self):
        a = " ".join(self._avisos_asistente("QCOM")).lower()
        self.assertIn("invalidaci", a)

    @falla(65, "hecho material del expediente sin citar en la tesis")
    def test_65_hecho_material_sin_citar(self):
        self.assertTrue(any("8-K" in a for a in self._avisos_asistente("QCOM")))

    @falla(73, "Kd por debajo del tipo libre de riesgo sin aviso")
    def test_73_kd_bajo_rf(self):
        self.assertTrue(any("Kd" in a or "coste de la deuda" in a for a in avisos("AAPL")))


if __name__ == "__main__":
    unittest.main()
