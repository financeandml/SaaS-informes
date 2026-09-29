"""El extractor: cada número de cada adjunto, con su rótulo, su periodo, su página y su rectángulo.

Un estado financiero en un PDF es texto con posición: una cabecera que dice qué
periodo es cada columna, una línea de escala («in thousands, except per share
data») y filas con un rótulo a la izquierda y tantas cifras como columnas. Aquí
se reconstruye eso a partir de las cajas de cada carácter, no del texto plano,
por dos razones que no son de comodidad:

- **La columna se asigna por posición horizontal**, no por orden de aparición.
  En el texto plano una celda vacía desplaza las siguientes una columna a la
  izquierda y la cifra de 2024 acaba atribuida a 2025 con la cifra cuadrando.
- **Cada candidato guarda su rectángulo en puntos**, que es lo que permite
  recortar la página exacta y adjuntar la imagen como evidencia (§7 de la
  plantilla). Sin coordenadas no hay recorte, y sin recorte la cifra viaja sin
  su prueba.

Lo que sale de aquí son *candidatos*: lecturas con su rótulo, la cabecera de
bloque que las precede («Earnings per share:»), el periodo inferido de la
cabecera y la escala de la página. Ninguno es un hecho hasta que `contraste`
lo pone al lado del hecho SEC o el analista lo confirma.

Para hojas de cálculo la mecánica es la misma con filas y columnas en lugar de
puntos, y la referencia es la celda («Income Statement!F8»).

Las cuentas españolas (PGC y NIIF, cuentas anuales y semestrales) se leen por
un camino propio (`_leer_es`) porque se imprimen de otra manera: cabeceras
«Notas 2025 2024» o «31/12/2025 31/12/2024*» sin el tipo de periodo, que se
deduce del estado (balance → instante) y del documento (anuales → 12 meses,
semestral → 6); punto de millar y coma decimal; columnas de notas, de códigos
del PGC y de casillas del modelo normalizado del Registro Mercantil, cuyos
números caen cerca de las cifras. Ahí manda la geometría: cada número va a la
columna cuyo centro tiene más cerca, y si lo más cercano es la columna de notas
o de casillas, no es un importe. El camino de EE. UU. no cambia.
"""

from __future__ import annotations

import calendar
import re
from dataclasses import dataclass, field
from datetime import date
from typing import Dict, Iterable, List, Optional, Sequence, Tuple

from .expediente import Adjunto, Tipo
from .hechos import Periodo

__all__ = ["Candidato", "Fila", "Linea", "PaginaLeida", "Token", "extraer_pdf", "extraer_xlsx", "leer_pagina"]

MESES = {m: i for i, m in enumerate(
    ["January", "February", "March", "April", "May", "June", "July",
     "August", "September", "October", "November", "December"], 1)}
_MES = "(?:January|February|March|April|May|June|July|August|September|October|November|December)"
_NUMERO = re.compile(r"^\(?-?\$?\s*\d[\d,.]*\)?%?$")
_ES_MILLARES = re.compile(r"^\d{1,3}(?:\.\d{3})+$")          # 10.542.801
_ES_DECIMAL = re.compile(r"^\d{1,3}(?:\.\d{3})*,\d+$")        # 0,68 · 1.234,5
_ES_INEQUIVOCO = re.compile(r"^\d{1,3}(?:\.\d{3})*,\d{1,2}$")  # «28,678» es ambiguo; «0,68» no
_EN_ESTANDAR = re.compile(r"^\d{1,3}(?:,\d{3})*(?:\.\d+)?$|^\d+(?:\.\d+)?$")
# formato inglés inequívoco: coma de millar y punto decimal a la vez, o dos grupos de millar con coma
_EN_INEQUIVOCO = re.compile(r"^\d{1,3}(?:,\d{3})+\.\d{1,2}$|^\d{1,3}(?:,\d{3}){2,}$")
_GUION = {"—", "–", "-", "— ", "—$"}


@dataclass
class Token:
    texto: str
    x0: float
    y0: float
    x1: float
    y1: float

    @property
    def xc(self) -> float:
        return (self.x0 + self.x1) / 2


@dataclass
class Linea:
    tokens: List[Token]

    @property
    def texto(self) -> str:
        return " ".join(t.texto for t in self.tokens)

    @property
    def y0(self) -> float:
        return min(t.y0 for t in self.tokens)

    @property
    def y1(self) -> float:
        return max(t.y1 for t in self.tokens)

    @property
    def x0(self) -> float:
        return min(t.x0 for t in self.tokens)

    @property
    def x1(self) -> float:
        return max(t.x1 for t in self.tokens)


@dataclass
class Columna:
    periodo: Optional[Periodo]
    xc: float                # centro horizontal (PDF) o índice de columna (XLSX)
    etiqueta: str            # lo que decía la cabecera, para citarla
    fin: Optional[date] = None


@dataclass
class Fila:
    rotulo: str
    contexto: str            # última línea de cabecera sin cifras encima de la fila
    celdas: Dict[int, "Candidato"]   # índice de columna → candidato
    rect: Tuple[float, float, float, float]


@dataclass
class Candidato:
    documento: str
    pagina: int              # índice físico del PDF o de la hoja (desde 1)
    rotulo: str
    contexto: str
    periodo: Optional[Periodo]
    etiqueta_columna: str
    valor: float             # ya escalado a unidades (USD, acciones), salvo si `por_accion`
    crudo: str
    escala: int
    por_accion: bool         # la escala no se aplica: la fila es «per share»
    guion: bool              # la celda era un guion: cero declarado por el documento
    porcentaje: bool
    rect: Tuple[float, float, float, float]     # del número (puntos, origen abajo-izq) o (fila, col, fila, col)
    rect_fila: Tuple[float, float, float, float]
    referencia: str = ""     # celda de hoja de cálculo
    # la moneda del documento, no la del mercado: la cifra y su unidad salen del mismo sitio (regla 13); "" si el
    # documento no la dice
    moneda: str = "USD"


@dataclass
class PaginaLeida:
    numero: int
    lineas: List[Linea]
    escala: int
    titulo: str                  # primera línea en mayúsculas o con «Consolidated…», si la hay
    columnas: List[Columna]
    filas: List[Fila]
    ancho: float = 0
    alto: float = 0
    # solo en cuentas españolas: qué estado es la página («balance», «resultados», «flujos», «patrimonio» o "" si es
    # una nota) y si es del grupo (True), de la sociedad (False) o no se sabe (None)
    estado: str = ""
    consolidado: Optional[bool] = None
    anclas: List[float] = field(default_factory=list)   # centros de las columnas de notas, códigos y casillas


# ---------------------------------------------------------------------------
# Del PDF a líneas y tokens
# ---------------------------------------------------------------------------

def _lineas_de(page, sin_girados: bool = False) -> List[Linea]:
    """Agrupa los caracteres en líneas (por altura) y en tokens (por hueco horizontal).

    Cada carácter se lee por su índice con `FPDFText_GetUnicode`, no del texto
    plano: el texto que devuelve `get_text_range` no comparte índice con las
    cajas de `get_charbox` (pdfium inserta y omite caracteres al linealizar), y
    con el texto plano salían rótulos sin espacios y cifras a las que les
    faltaba el primer grupo de dígitos.

    `sin_girados` deja fuera el texto girado cuando es minoría en la página: el código de verificación de la firma
    electrónica que las cuentas españolas imprimen en vertical en el margen caía, letra a letra, en cada renglón de
    la tabla («f ACTIVO CORRIENTE», un «7» junto a cada cifra).
    """
    import math

    import pypdfium2.raw as pdfium_c
    tp = page.get_textpage()
    n = tp.count_chars()
    girados = set()
    if sin_girados:
        for i in range(n):
            angulo = pdfium_c.FPDFText_GetCharAngle(tp, i)
            # −1 es «no se sabe»: se trata como horizontal
            if angulo >= 0 and min(angulo, 2 * math.pi - angulo) > 0.05:
                girados.add(i)
        if len(girados) * 2 >= n:
            girados = set()      # una página girada entera no es un margen: se deja como estaba
    chars = []
    espacio = False
    for i in range(n):
        if i in girados:
            continue
        codigo = pdfium_c.FPDFText_GetUnicode(tp, i)
        ch = chr(codigo) if codigo else ""
        if not ch or ch in "\r\n":
            continue
        if ch.isspace():
            # El glifo de espacio tiene caja de anchura cero y a veces cae, por
            # redondeo, a la derecha del carácter que lo sigue («March 3 1,»).
            # No se guarda por su posición, pero sí se recuerda que estaba: la
            # capa de texto sabe dónde separan las palabras aunque el hueco no lo
            # diga, y en las cursivas en negrita de los epígrafes del Item 1A el
            # hueco entre palabras es menor que el criterio de abajo («ofour»).
            espacio = True
            continue
        x0, y0, x1, y1 = tp.get_charbox(i, loose=True)
        if x1 - x0 <= 0 or y1 - y0 <= 0:
            continue
        chars.append((ch, x0, y0, x1, y1, espacio))
        espacio = False
    chars.sort(key=lambda c: (-round((c[2] + c[4]) / 2, 0), c[1]))
    lineas: List[List[tuple]] = []
    for c in chars:
        yc = (c[2] + c[4]) / 2
        if lineas and abs(((lineas[-1][0][2] + lineas[-1][0][4]) / 2) - yc) <= 2.5:
            lineas[-1].append(c)
        else:
            lineas.append([c])
    salida = []
    for l in lineas:
        l.sort(key=lambda c: c[1])
        tokens = _tokens_de(l)
        if tokens:
            salida.append(Linea(tokens))
    tp.close()       # las líneas ya son texto y coordenadas: sin cerrarlo, pdfium se queda con la página y el fichero
    return salida


def _tokens_de(caracteres) -> List[Token]:
    """Los tokens de una línea a partir de sus caracteres: (letra, x0, y0, x1, y1, espacio antes)."""
    tokens: List[Token] = []
    for ch, x0, y0, x1, y1, espacio_antes in caracteres:
        altura = max(y1 - y0, 1)
        hueco = (x0 - tokens[-1].x1) if tokens else 0
        # Las palabras se separan por posición: un hueco de más de 0,12 em
        # dentro del token es un espacio; uno de más de 1,2 em cierra el token
        # (es un hueco de columna). El «$» es siempre token aparte: va pegado
        # al número de la columna anterior y, fundido con él, la cifra dejaba
        # de ser un número.
        es_dolar = ch == "$"
        if tokens and hueco < altura * 1.2 and not es_dolar and tokens[-1].texto != "$":
            t = tokens[-1]
            if hueco > altura * 0.12 or espacio_antes:
                t.texto += " "
            t.texto += ch; t.x1 = max(t.x1, x1); t.y0 = min(t.y0, y0); t.y1 = max(t.y1, y1)
        else:
            tokens.append(Token(ch, x0, y0, x1, y1))
    return [Token(t.texto.strip(), t.x0, t.y0, t.x1, t.y1) for t in tokens if t.texto.strip() and t.texto.strip() != "$"]


def _es_numero(t: str) -> bool:
    # el «€» va pegado a la cifra en algunas cuentas españolas («1.856.554,35 €»): se quita como el «$»
    t = t.replace(" ", "").replace("€", "")
    return bool(_NUMERO.match(t)) or t in _GUION


def _valor_de(t: str, locale_es: bool = False) -> Tuple[Optional[float], bool, bool]:
    """(valor, es_guion, es_porcentaje). Un guion es cero declarado por el documento.

    `locale_es`: la página usa punto de millar y coma decimal («10.542.801»,
    «0,68»), como un PDF impreso desde Excel en español. Se decide por página,
    no por token: «1.234» solo es mil doscientos treinta y cuatro si el resto
    de la página lo es.
    """
    s = t.replace(" ", "").replace("€", "")
    if s in _GUION:
        return 0.0, True, False
    pct = s.endswith("%") or s.endswith("%)")
    neg = s.startswith("(") and s.endswith(")")
    s = s.strip("()%$").replace("$", "")
    if s.startswith("-"):
        neg, s = True, s[1:]
    if not s:
        return None, False, False
    if locale_es and (_ES_MILLARES.match(s) or _ES_DECIMAL.match(s)):
        s = s.replace(".", "").replace(",", ".")
    elif _EN_ESTANDAR.match(s):
        s = s.replace(",", "")
    elif _ES_MILLARES.match(s) and s.count(".") >= 2:
        s = s.replace(".", "")
    else:
        return None, False, False
    try:
        v = float(s)
    except ValueError:
        return None, False, False
    return (-v if neg else v), False, pct


def _locale_es(lineas: Sequence[Linea]) -> bool:
    """Coma decimal o dos grupos de millar con punto en algún token: la página está en español."""
    for l in lineas:
        for tok in l.tokens:
            s = tok.texto.replace(" ", "").strip("()$%")
            if _ES_INEQUIVOCO.match(s) or (_ES_MILLARES.match(s) and s.count(".") >= 2):
                return True
    return False


def _locale_en(lineas: Sequence[Linea]) -> bool:
    """Coma de millar y punto decimal en algún token: la página está en formato inglés aunque el documento sea
    español (hay semestrales que anexan estados impresos así, «1,914,716.70 €»)."""
    return any(_EN_INEQUIVOCO.match(tok.texto.replace(" ", "").replace("€", "").strip("()$%-"))
               for l in lineas for tok in l.tokens)


def _escala_de(lineas: Sequence[Linea], hasta: Optional[int] = None, espanol: bool = False, por_defecto: int = 1) -> int:
    """La escala declarada en la página («in thousands»), tomando la última declaración
    que precede a la cabecera de la tabla (`hasta`) o, si no hay ninguna antes, la
    primera de la página. En un estado va en la cabecera; en una nota va donde el
    redactor la puso, a veces 17 líneas más abajo, y con solo doce líneas la nota de
    ingresos por región salía sin escala y las regiones no cuadraban con el total.

    En un documento español la declaración es «(Expresados en euros)», «Miles de euros» o «en millones de euros»;
    solo cuenta en una línea corta —la de la cabecera o la del título—: en la prosa de una nota, «19,4 miles de
    euros» es una cifra, no la escala de la tabla."""
    def escala_en(t: str) -> int:
        corta = len(t) <= 80
        t = t.lower().replace(" ", "")
        if "inthousands" in t:
            return 1_000
        if "inmillions" in t:
            return 1_000_000
        if "inbillions" in t:
            return 1_000_000_000
        if espanol and corta:
            if re.search(r"milesdeeuros|enmiles\b|\(miles\)|milesde€|miles€|k€", t):
                return 1_000
            if re.search(r"millonesdeeuros|enmillones|millonesde€|m€\)", t):
                return 1_000_000
            if re.search(r"(?:expresad[oa]s?|cifras|importes)en(?:unidadesde)?(?:euros|€)|\(euros\)|eneuros\)|^euros$", t):
                return 1
        return 0
    declaraciones = [(k, escala_en(l.texto)) for k, l in enumerate(lineas) if escala_en(l.texto)]
    if not declaraciones:
        return por_defecto
    if hasta is not None:
        previas = [e for k, e in declaraciones if k < hasta]
        if previas:
            return previas[-1]
    return declaraciones[0][1]


_TITULO_ES = re.compile(r"(?i)\bbalance\b|situaci[óo]n financiera|p[ée]rdidas y ganancias|cuenta de (?:resultados|explotaci[óo]n)"
                        r"|flujos? de efectivo|ingresos y gastos reconocidos|cambios en el patrimonio|resultado global")


def _titulo_de(lineas: Sequence[Linea], espanol: bool = False) -> str:
    for l in lineas[:6]:
        t = l.texto.strip()
        if re.search(r"(?i)consolidated ?(statements?|balance)", t):
            return t
        if espanol and _TITULO_ES.search(t):
            return t
    return lineas[0].texto.strip() if lineas else ""


# ---------------------------------------------------------------------------
# Cabeceras de columna
# ---------------------------------------------------------------------------

def _fin_mes(anio: int, mes: int) -> date:
    return date(anio, mes, calendar.monthrange(anio, mes)[1])


def _es_por_accion(bloque: str) -> bool:
    """Si la fila es un importe por accion (no se escala) o no.

    «Shares used in per share calculations» dice «per share» y sin embargo es un recuento de acciones: lleva su
    escala, 1.096 son 1.096 millones de acciones. Sin esta salvedad las acciones medias se leian sin escalar y su
    fila «Basic» se colaba como BPA, discrepando de la SEC por mil millones.
    """
    if re.search(r"(?i)shares used|weighted[- ]average (?:number of )?shares|shares outstanding", bloque):
        return False
    return bool(re.search(r"(?i)per share|per common share", bloque))


def _cierre(anio: int, mes: int, dia: Optional[int]) -> date:
    """El cierre que declara la cabecera: «June 28, 2026» cierra el 28, no el 30.

    Quien cierra por semanas (52/53) no cierra a fin de mes, y redondear al último día dejaba cada columna del
    documento fechada dos días después que el hecho de la SEC: ni una cifra leída casaba con su fuente.
    """
    if dia and 1 <= dia <= calendar.monthrange(anio, mes)[1]:
        return date(anio, mes, dia)
    return _fin_mes(anio, mes)


def _columnas_de(lineas: Sequence[Linea]) -> Tuple[List[Columna], int]:
    """Lee la cabecera: segmentos de tipo («Three Months Ended», «Year ended», «As of»),
    tokens de mes («June 30,») y tokens de año, y los casa por posición horizontal.
    Devuelve las columnas y el índice de la línea de años (las filas vienen después)."""
    tipos: List[Tuple[str, float, float, Optional[int]]] = []   # (texto, x0, x1, mes embebido)
    meses: List[Tuple[int, int, float]] = []                     # (mes, día, xc)
    anios: List[Tuple[int, float]] = []
    idx_anios = -1
    for i, l in enumerate(lineas[:25]):
        t = l.texto
        m_tipo = re.finditer(r"((?:Three|Six|Nine|Twelve) Months(?: Ended)?|Year[s]? [Ee]nded|As of|Fiscal Year[s]? Ended|Quarter Ended)(?:\s+(" + _MES + r")\s+(\d{1,2}),?)?", t)
        desde = 0
        for m in m_tipo:
            # posición del segmento: los tokens de la línea cuyo texto compone el segmento;
            # dos «Three Months Ended» en la misma línea se buscan cada uno a partir del anterior
            x0, x1, desde = _rango_de(l, m.group(0), desde)
            mes = MESES[m.group(2)] if m.group(2) else None
            tipos.append((m.group(1), x0, x1, mes, int(m.group(3)) if m.group(3) else None))
        for tok in l.tokens:
            # los tres cierres de la cabecera pueden venir en un solo token —«September 28, September 29, September
            # 24,»—: se buscan dentro del texto y el centro de cada uno se reparte a lo ancho del token, que es lo
            # único que hace falta para casarlos con sus años (van en el mismo orden).
            ancho = max(tok.x1 - tok.x0, 1e-6)
            largo = max(len(tok.texto), 1)
            for mm in re.finditer(r"(" + _MES + r")\s+(\d{1,2}),?", tok.texto):
                centro = tok.x0 + ancho * (mm.start() + mm.end()) / (2 * largo)
                meses.append((MESES[mm.group(1)], int(mm.group(2)), centro))
        toks_anio = [tok for tok in l.tokens if re.fullmatch(r"(?:19|20)\d{2}", tok.texto.strip())]
        # A la izquierda del primer año puede ir el rótulo de la columna de conceptos —«(in millions, except per
        # share data)»— y a la derecha solo pueden ir años y meses: así es como se imprime la cabecera de un estado
        # financiero. Exigir que la línea entera fueran años dejaba fuera el balance y la cuenta de resultados de
        # cualquier emisor que ponga las unidades en esa misma línea.
        k0 = next((k for k, tok in enumerate(l.tokens) if re.fullmatch(r"(?:19|20)\d{2}", tok.texto.strip())), None)
        resto_cabecera = k0 is not None and all(
            re.fullmatch(r"(?:19|20)\d{2}|" + _MES + r"\s+\d{1,2},?|As of|Ended|and", tok.texto.strip()) for tok in l.tokens[k0:])
        rotulo = " ".join(tok.texto for tok in l.tokens[:k0]) if k0 else ""
        # el rótulo no puede traer importes: «Vencimientos 1,250 2026 2027» es una fila, no una cabecera
        rotulo_limpio = not re.search(r"\d{1,3}(?:,\d{3})+|\d+\.\d", rotulo)
        if len(toks_anio) >= 2 and resto_cabecera and rotulo_limpio and idx_anios < 0:
            anios = [(int(tok.texto), tok.xc) for tok in toks_anio]
            idx_anios = i
            break
    if not anios:
        return [], -1
    columnas = []
    reparto = _repartir(tipos, [xc for _, xc in anios])
    for k, (anio, xc) in enumerate(anios):
        tipo = reparto[k]
        # El mes: del token de mes de la columna (uno por columna), del único
        # token de mes de la cabecera (centrado sobre todas), o del propio
        # segmento («Year ended December 31,»). Si no hay ninguno, no se inventa:
        # la columna queda sin periodo y la página sin filas.
        if len(meses) == len(anios):
            mes, dia = sorted(meses, key=lambda m: m[2])[k][:2]
        elif len(meses) == 1:
            mes, dia = meses[0][:2]
        elif meses and abs(min(meses, key=lambda m: abs(m[2] - xc))[2] - xc) < 60:
            mes, dia = min(meses, key=lambda m: abs(m[2] - xc))[:2]
        else:
            mes, dia = (tipo[3], tipo[4]) if tipo else (None, None)
        if mes is None:
            columnas.append(Columna(periodo=None, xc=xc, etiqueta=f"{anio} (sin mes en la cabecera)", fin=None))
            continue
        fin = _cierre(anio, mes, dia)
        etiqueta = f"{tipo[0] if tipo else ''} {calendar.month_name[mes]} {fin.day}, {anio}".strip()
        texto_tipo = (tipo[0] if tipo else "").lower()
        if "three" in texto_tipo or "quarter" in texto_tipo:
            p = Periodo.de_meses(fin, 3)
        elif "six" in texto_tipo:
            p = Periodo.de_meses(fin, 6)
        elif "nine" in texto_tipo:
            p = Periodo.de_meses(fin, 9)
        elif "twelve" in texto_tipo or "year" in texto_tipo:
            p = Periodo.de_meses(fin, 12)
        elif "as of" in texto_tipo:
            p = Periodo.instante(fin)
        else:
            p = None   # sin tipo: un balance sin «As of» se resuelve por el título de la página
        columnas.append(Columna(periodo=p, xc=xc, etiqueta=etiqueta, fin=fin))
    return columnas, idx_anios


def _rango_de(linea: Linea, fragmento: str, desde: int = 0) -> Tuple[float, float, int]:
    """El tramo horizontal que ocupa `fragmento` en la línea, palabra a palabra.

    Un token puede contener varias palabras («Three Months Ended»), así que la
    búsqueda se hace sobre la lista de palabras con el token del que sale cada
    una; buscando por tokens enteros, «Three Months Ended» se extendía hasta el
    final de «Six Months Ended» y las columnas de seis meses salían como trimestre.
    """
    palabras = fragmento.split()
    secuencia = [(w, t) for t in linea.tokens for w in t.texto.split()]
    for i in range(desde, len(secuencia) - len(palabras) + 1):
        if [w for w, _ in secuencia[i:i + len(palabras)]] == palabras:
            return secuencia[i][1].x0, secuencia[i + len(palabras) - 1][1].x1, i + len(palabras)
    return linea.x0, linea.x1, desde


def _repartir(tipos, xcs):
    """Qué segmento de cabecera cubre cada columna.

    Cada segmento («Three Months Ended», «Twelve Months Ended») va centrado sobre
    su grupo de columnas, pero su texto no ocupa el ancho del grupo: uno de cuatro
    columnas mide lo mismo que uno de una. Así que ni «el segmento cuyo tramo
    contiene la columna» ni «el de centro más cercano» aciertan cuando un grupo
    ancho está junto a uno estrecho. Se prueba cada partición contigua de las
    columnas en tantos grupos como segmentos (en orden) y se queda la que
    mejor alinea el centro de cada grupo con el centro de su segmento.
    """
    if not tipos or not xcs:
        return [None] * len(xcs)
    orden = sorted(tipos, key=lambda t: t[1])
    n, k = len(xcs), len(orden)
    if k > n:
        orden = orden[:n]; k = n
    if k == 1:
        return [orden[0]] * n
    from itertools import combinations
    centros = [(t[1] + t[2]) / 2 for t in orden]
    mejor, mejor_coste = None, None
    for cortes in combinations(range(1, n), k - 1):
        limites = (0,) + cortes + (n,)
        coste = 0.0
        for g in range(k):
            grupo = xcs[limites[g]:limites[g + 1]]
            coste += abs((grupo[0] + grupo[-1]) / 2 - centros[g])
        if mejor_coste is None or coste < mejor_coste:
            mejor, mejor_coste = limites, coste
    asignacion = []
    for g in range(k):
        asignacion += [orden[g]] * (mejor[g + 1] - mejor[g])
    return asignacion


# ---------------------------------------------------------------------------
# Filas y candidatos
# ---------------------------------------------------------------------------

def leer_pagina(page, numero: int, documento: str, espanol: bool = False, cierre: Optional[date] = None,
                meses: Optional[int] = None, moneda: str = "USD", previa: Optional["PaginaLeida"] = None) -> PaginaLeida:
    """Una página de estados. Con `espanol`, la de unas cuentas españolas: `cierre` es el del documento (el año de la
    cabecera «2025» lleva su día y su mes), `meses` la duración de sus columnas de flujo (12 en las anuales, 6 en el
    semestral) y `previa` la página anterior, cuya cabecera hereda una continuación que no la repite."""
    if espanol:
        ancho, alto = page.get_size()
        return _leer_es(_lineas_de(page, sin_girados=True), numero, documento, ancho, alto, cierre, meses, moneda, previa)
    lineas = _lineas_de(page)
    ancho, alto = page.get_size()
    titulo = _titulo_de(lineas)
    columnas, idx = _columnas_de(lineas)
    escala = _escala_de(lineas, idx if idx >= 0 else None)
    # Un balance sin «As of» en la cabecera (el 10-K lo lleva; la web del emisor no):
    # el título de la página dice que son instantes.
    if columnas and re.search(r"(?i)balance sheet", titulo):
        for c in columnas:
            if c.periodo is None and c.fin is not None:
                c.periodo = Periodo.instante(c.fin)
    locale_es = _locale_es(lineas)
    filas: List[Fila] = []
    contexto = ""
    for l in lineas[idx + 1:] if idx >= 0 else []:
        toks = l.tokens
        numericos = [t for t in toks if _es_numero(t.texto) and t.texto.strip() not in ("$",)]
        etiqueta = [t for t in toks if t not in numericos and t.texto.strip() != "$"]
        rotulo = " ".join(t.texto for t in etiqueta).strip(" :")
        if not numericos:
            if rotulo:
                contexto = rotulo
            continue
        if not rotulo:
            continue
        es_por_accion = _es_por_accion(contexto + " " + rotulo)
        celdas: Dict[int, Candidato] = {}
        for t in numericos:
            v, guion, pct = _valor_de(t.texto, locale_es)
            if v is None:
                continue
            j = _columna_de(columnas, t.xc)
            if j is None:
                continue
            col = columnas[j]
            valor = v if (es_por_accion or pct) else v * escala
            celdas[j] = Candidato(documento=documento, pagina=numero, rotulo=rotulo, contexto=contexto,
                                  periodo=col.periodo, etiqueta_columna=col.etiqueta, valor=valor, crudo=t.texto,
                                  escala=escala, por_accion=es_por_accion, guion=guion, porcentaje=pct,
                                  rect=(t.x0, t.y0, t.x1, t.y1), rect_fila=(l.x0, l.y0, l.x1, l.y1))
        if celdas:
            filas.append(Fila(rotulo=rotulo, contexto=contexto, celdas=celdas, rect=(l.x0, l.y0, l.x1, l.y1)))
    return PaginaLeida(numero=numero, lineas=lineas, escala=escala, titulo=titulo, columnas=columnas,
                       filas=filas, ancho=ancho, alto=alto)


def _columna_de(columnas: Sequence[Columna], xc: float) -> Optional[int]:
    """La columna cuyo centro está más cerca; un número alineado a la derecha cae algo a la derecha del centro."""
    if not columnas:
        return None
    j = min(range(len(columnas)), key=lambda k: abs(columnas[k].xc - xc))
    if abs(columnas[j].xc - xc) > 70:
        return None
    return j


def extraer_pdf(adjunto: Adjunto, paginas: Optional[Iterable[int]] = None) -> List[PaginaLeida]:
    """Lee las páginas con tablas (cabecera de años + filas con cifras). Solo PDF.

    En unas cuentas españolas, además, cada página sabe si es del grupo o de la sociedad: lo dice ella misma («Y
    SOCIEDADES DEPENDIENTES», «Empresa: … S.A.») o lo dijo la portada de su sección —el informe de auditoría de las
    cuentas consolidadas, el de las individuales—, que es lo único que distingue dos balances con el mismo modelo."""
    import pypdfium2 as pdfium
    espanol = adjunto.tipo in (Tipo.CCAA, Tipo.SEMESTRAL)
    moneda = moneda_de(adjunto) if espanol else "USD"
    meses = 6 if adjunto.tipo is Tipo.SEMESTRAL else 12
    doc = pdfium.PdfDocument(str(adjunto.ruta))
    try:
        indices = list(paginas) if paginas is not None else range(1, len(doc) + 1)
        salida = []
        ambito: Optional[bool] = None
        previa: Optional[PaginaLeida] = None
        for n in indices:
            texto = adjunto.paginas[n - 1] if n - 1 < len(adjunto.paginas) else ""
            if espanol:
                seccion = _ambito_de_seccion(texto)
                if seccion is not None:
                    ambito = seccion
                # la continuación de un estado no repite la cabecera ni, a veces, ningún año: basta con que traiga cifras
                if not re.search(r"\d{1,3}(?:[.,]\d{3})+|\d+,\d+", texto):
                    previa = None
                    continue
            elif not re.search(r"\b(19|20)\d{2}\b", texto) or not re.search(r"\d{1,3}(,\d{3})+|\d+\.\d+", texto):
                continue
            pagina = doc[n - 1]
            if espanol:
                p = leer_pagina(pagina, n, adjunto.nombre, espanol=True, cierre=adjunto.periodo_fin, meses=meses,
                                moneda=moneda, previa=previa if previa is not None and previa.numero == n - 1 else None)
                if p.consolidado is not None:
                    ambito = p.consolidado
                p.consolidado = ambito
                previa = p if p.columnas and p.estado in ("balance", "resultados", "flujos") else None
            else:
                p = leer_pagina(pagina, n, adjunto.nombre)
            pagina.close()
            if p.columnas and p.filas:
                salida.append(p)
    finally:
        doc.close()      # en Windows un PDF que no se cierra queda bloqueado: el analista no podría ni quitarlo
    return salida


# ---------------------------------------------------------------------------
# Cuentas españolas (PGC y NIIF): anuales y semestrales
# ---------------------------------------------------------------------------

MESES_ES = {m: i for i, m in enumerate(["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto",
                                        "septiembre", "octubre", "noviembre", "diciembre"], 1)}
MESES_ES["setiembre"] = 9
# el año puede venir partido por la capa de texto («31 de diciembre de 202 5»)
FECHA_ES = re.compile(r"(\d{1,2})\s+de\s+(" + "|".join(MESES_ES) + r")\s+(?:de|del)\s+((?:\d\s?){3}\d)\b", re.I)
_ANIO_CAB = re.compile(r"^((?:19|20)\d{2})\*?$")
_FECHA_CAB = re.compile(r"^(\d{1,2})[/.](\d{1,2})[/.]((?:19|20)\d{2})\*?$")
# columnas de la cabecera que no son importes: sus números —notas de la memoria, códigos del PGC, casillas del modelo
# del Registro Mercantil— caen cerca de las cifras, y leídos como importes serían cifras inventadas
_ANCLA = re.compile(r"(?i)^(?:notas?|n\.?º|casillas?|is|l[íi]nea|c[óo]digo|ref\.?|referencia)$")
# lo que acompaña a los años en la cabecera sin ser una columna
_ACOMPANA = re.compile(r"(?i)^(?:eur|euros|€|\(\*\)|\*)$")
# una referencia a notas de la memoria: «6», «9 y 15», «18.e», «9, 9.a, 19», «19 (c)», «4,13,23», «14,»
_REFERENCIA_NOTA = re.compile(r"^\(?\d{1,2}(?:\.\w{1,3})?(?:\s*\(\w\))?\)?"
                              r"(?:\s*(?:,|y|e|-)\s*\d{1,2}(?:\.\w{1,3})?(?:\s*\(\w\))?)*,?$")
# prefijos de partida del PGC («A)», «A-1)», «A.1)», «1.», «a)», «a1)», «VII.») y sufijos de fórmula («(A+B)»,
# «(1+2+…+13)», «( A.4 + 21 )», «(del 1 al 13)»): no son parte del nombre de la partida
_PREFIJO = re.compile(r"^(?:[A-Z](?:[.-]\d{1,2})?\)|\d{1,2}\.(?!\d)|[a-z]\d?\)|[IVXL]{1,5}\.)\s*")
_FORMULA = re.compile(r"\s*\(\s*[A-Z0-9][A-Z0-9.]*(?:\s*[+\-−]\s*[A-Z0-9][A-Z0-9.]*)+\s*\)\s*$")
_FORMULA_DEL = re.compile(r"\s*\((?:del|de la) \d+ al \d+\)\s*$", re.I)
# un rótulo que acaba en preposición, artículo o conjunción sigue en la línea de abajo
_CONTINUA = re.compile(r"(?i)\b(?:de|del|la|las|los|el|al|a|y|e|o|u|en|por|para|con|sobre|entre)$")
# filas que abren una sección del balance aunque traigan su subtotal: el contexto de lo que va debajo
_ENCABEZADO = re.compile(r"(?i)^(?:total )?(?:activos?|pasivos?) (?:no )?corrientes?$|^patrimonio neto$"
                         r"|^deudas a (?:largo|corto) plazo$")
_TITULO_SECCION = re.compile(r"(?i)informe de auditor[íi]a (?:independiente )?de (?:las )?cuentas anuales"
                             r"|informe de revisi[óo]n limitada|estados financieros intermedios")
_TITULO_CUENTAS = re.compile(r"(?i)cuentas anuales(?: consolidadas)? (?:al|a) \d{1,2} de ")
_DEL_GRUPO = re.compile(r"(?i)consolidad|sociedades dependientes|\bgrupo\b")
_INDIVIDUAL = re.compile(r"(?i)\bindividual(?:es)?\b")
_SOCIEDAD = re.compile(r"^(?:Empresa:\s*\d*\s*)?[A-ZÁÉÍÓÚÑ0-9][A-ZÁÉÍÓÚÑ0-9 .,&'-]{2,},?\s+S\.?\s?[AL]\.?(?:U\.?)?$")


def fecha_es(texto: str) -> Optional[date]:
    """La primera fecha larga en español del texto («31 de diciembre de 2025»)."""
    m = FECHA_ES.search(texto or "")
    if not m:
        return None
    try:
        return date(int(m.group(3).replace(" ", "")), MESES_ES[m.group(2).lower()], int(m.group(1)))
    except ValueError:
        return None


def limpiar_rotulo_es(r: str) -> str:
    """El nombre de la partida sin su prefijo del PGC ni su fórmula: «A.1) RESULTADO DE EXPLOTACIÓN (1+2+…+13)» es
    «RESULTADO DE EXPLOTACIÓN»."""
    r = _PREFIJO.sub("", r.strip())
    r = _FORMULA_DEL.sub("", _FORMULA.sub("", r))
    return re.sub(r"\s+", " ", r).strip(" :")


def moneda_de(adjunto: Adjunto) -> str:
    """La moneda en que el documento dice que expresa sus cifras; "" si no lo dice."""
    texto = " ".join(adjunto.paginas)
    eur = len(re.findall(r"(?i)\beuros?\b|€|\bEUR\b", texto))
    usd = len(re.findall(r"(?i)d[óo]lares|\bUSD\b|US\$", texto))
    if not eur and not usd:
        return ""
    return "EUR" if eur >= usd else "USD"


def _ambito_de_seccion(texto: str) -> Optional[bool]:
    """Si la página abre una sección del documento —el informe de auditoría o de revisión, la portada de las cuentas—,
    de quién es lo que sigue: del grupo (True) o de la sociedad (False). None si no abre ninguna."""
    cabeza = " ".join((texto or "").split())[:700]
    m = _TITULO_SECCION.search(cabeza) or _TITULO_CUENTAS.search(cabeza[:250])
    if not m:
        return None
    trozo = cabeza[max(0, m.start() - 150):m.end() + 120]
    if _INDIVIDUAL.search(trozo):
        return False
    return bool(_DEL_GRUPO.search(trozo))


def _ambito_de_pagina(titulo: Sequence[Linea]) -> Optional[bool]:
    """Lo que dice de sí mismo el título de un estado: «… Y SOCIEDADES DEPENDIENTES» o «consolidado» es del grupo, el
    nombre de la sociedad a secas es de la sociedad. Solo las líneas de encima de la cabecera: más abajo, «empresas
    del grupo» es una partida, no un título."""
    for l in titulo[:6]:
        if _INDIVIDUAL.search(l.texto):
            return False
        if _DEL_GRUPO.search(l.texto):
            return True
    if any(_SOCIEDAD.match(l.texto.strip()) for l in titulo[:4]):
        return False
    return None


def _partir(tok: Token) -> List[Token]:
    """Las palabras de un token con su tramo horizontal, repartido a lo ancho por su posición en el texto: la cabecera
    «31/12/2024 Casilla IS» llega como un solo token, y la fecha y la casilla son dos columnas distintas."""
    palabras = list(re.finditer(r"\S+", tok.texto))
    if len(palabras) <= 1:
        return [tok]
    ancho, largo = max(tok.x1 - tok.x0, 1e-6), max(len(tok.texto), 1)
    return [Token(m.group(0), tok.x0 + ancho * m.start() / largo, tok.y0, tok.x0 + ancho * m.end() / largo, tok.y1)
            for m in palabras]


def _es_fecha_cab(t: str) -> bool:
    return bool(_ANIO_CAB.match(t) or _FECHA_CAB.match(t))


def _columnas_es(lineas: Sequence[Linea], cierre: Optional[date]) -> Tuple[List[Columna], int, List[float]]:
    """La cabecera de un estado español: «Notas 2025 2024», «Nota 31/12/2025 31/12/2024*» o «Línea Descripción Notas
    30/06/2025 31/12/2024 Casilla IS». El año solo trae su día y su mes del cierre del documento; el tipo de periodo
    no lo dice la cabecera, lo dice el estado (balance → instante) y el documento (anuales → 12 meses). Devuelve las
    columnas sin periodo, el índice de la línea de cabecera y los centros de las columnas que no son importes."""
    for i, l in enumerate(lineas[:30]):
        palabras = [w for t in l.tokens for w in _partir(t)]
        k0 = next((k for k, w in enumerate(palabras) if _es_fecha_cab(w.texto)), None)
        if k0 is None:
            continue
        cola = palabras[k0:]
        fechas = [w for w in cola if _es_fecha_cab(w.texto)]
        if len(fechas) < 2 or not all(_es_fecha_cab(w.texto) or _ANCLA.match(w.texto) or _ACOMPANA.match(w.texto)
                                      for w in cola):
            continue
        # el rótulo no puede traer importes: «Vencimientos 1.250 2026 2027» es una fila, no una cabecera
        if re.search(r"\d{1,3}(?:[.,]\d{3})+|\d+,\d", " ".join(w.texto for w in palabras[:k0])):
            continue
        columnas = []
        for w in fechas:
            m = _FECHA_CAB.match(w.texto)
            try:
                if m:
                    fin = date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
                elif cierre is not None:
                    fin = _cierre(int(_ANIO_CAB.match(w.texto).group(1)), cierre.month, cierre.day)
                else:
                    fin = None       # sin cierre del documento ni fecha en la cabecera, el año solo no dice qué día
            except ValueError:
                fin = None
            columnas.append(Columna(periodo=None, xc=w.xc, etiqueta=w.texto, fin=fin))
        return columnas, i, [w.xc for w in palabras if _ANCLA.match(w.texto)]
    return [], -1, []


def _estado_de(lineas: Sequence[Linea], idx: int) -> str:
    """Qué estado es la página, por su título y su cabecera. Una nota no es ninguno ("")."""
    cabeza = " ".join(l.texto for l in lineas[:idx + 1])
    if re.search(r"(?i)ingresos y gastos reconocidos|cambios en el patrimonio|resultado global|resultado integral", cabeza):
        return "patrimonio"
    if re.search(r"(?i)flujos? de (?:efectivo|caja)", cabeza):
        return "flujos"
    if re.search(r"(?i)p[ée]rdidas y ganancias|cuenta de (?:resultados|explotaci[óo]n)|\(debe\)\s*/\s*haber"
                 r"|resultados? (?:intermedi|consolidad)", cabeza):
        return "resultados"
    if re.search(r"(?i)\bbalance\b|situaci[óo]n financiera", cabeza) \
            or re.match(r"(?i)\s*(?:activo|pasivo|patrimonio neto)\b", lineas[idx].texto if idx >= 0 else ""):
        return "balance"
    return ""


def _destino(columnas: Sequence[Columna], anclas: Sequence[float], xc: float) -> Optional[int]:
    """La columna de importes a la que va un número, o None si lo más cercano es una columna de notas, de códigos o de
    casillas, o si no cae cerca de ninguna."""
    mejor, distancia = None, None
    for k, c in enumerate(columnas):
        if distancia is None or abs(c.xc - xc) < distancia:
            mejor, distancia = k, abs(c.xc - xc)
    for a in anclas:
        if distancia is None or abs(a - xc) < distancia:
            mejor, distancia = None, abs(a - xc)
    if mejor is None or distancia > 70:
        return None
    return mejor


def _sin_notas(lineas: Sequence[Linea], columnas: Sequence[Columna], anclas: Sequence[float]) -> List[Linea]:
    """Fuera los números y las referencias a notas que caen en las columnas de notas, códigos o casillas: ni son
    importes ni son parte del nombre de la partida («Importe neto de la cifra de negocios 18.a»)."""
    if not anclas:
        return list(lineas)
    salida = []
    for l in lineas:
        toks = [t for t in l.tokens
                if not ((_es_numero(t.texto) or _REFERENCIA_NOTA.match(t.texto.strip()))
                        and min(abs(a - t.xc) for a in anclas) < min((abs(c.xc - t.xc) for c in columnas), default=1e9))]
        if toks:
            salida.append(Linea(toks))
    return salida


def _solo_cifras(l: Linea) -> bool:
    return all(_es_numero(t.texto) for t in l.tokens)


def _fundir(lineas: Sequence[Linea]) -> List[Linea]:
    """Un rótulo en negrita y sus cifras en redonda comparten renglón pero no altura exacta, y la capa de texto los da
    como dos líneas: «TOTAL ACTIVO» por un lado y «10.314.343,63 11.763.690,19» por otro. Se juntan cuando una es
    solo rótulo, la otra solo cifras y se solapan en vertical más de media altura."""
    salida: List[Linea] = []
    for l in lineas:
        if salida:
            a = salida[-1]
            rotulo_y_cifras = (not any(_es_numero(t.texto) for t in a.tokens) and _solo_cifras(l)) or \
                              (_solo_cifras(a) and not any(_es_numero(t.texto) for t in l.tokens))
            solape = min(a.y1, l.y1) - max(a.y0, l.y0)
            if rotulo_y_cifras and solape >= 0.5 * min(a.y1 - a.y0, l.y1 - l.y0):
                salida[-1] = Linea(sorted(a.tokens + l.tokens, key=lambda t: t.x0))
                continue
        salida.append(l)
    return salida


def _alinea(lineas: Sequence[Linea], columnas: Sequence[Columna], anclas: Sequence[float]) -> bool:
    """Si las primeras líneas de una página sin cabecera traen cifras en las columnas de la página anterior: es la
    continuación del mismo estado. Una nota que empieza con prosa no lo es."""
    for l in lineas[:4]:
        for t in l.tokens:
            v, _, _ = _valor_de(t.texto, True)
            if v is not None and _destino(columnas, anclas, t.xc) is not None:
                return True
    return False


def _leer_es(lineas: List[Linea], numero: int, documento: str, ancho: float, alto: float, cierre: Optional[date],
             meses: Optional[int], moneda: str, previa: Optional[PaginaLeida] = None) -> PaginaLeida:
    """Las filas de una página de unas cuentas españolas (ver `leer_pagina`)."""
    titulo = _titulo_de(lineas, espanol=True)
    cierre = cierre or fecha_es(" ".join(l.texto for l in lineas[:8]))
    columnas, idx, anclas = _columnas_es(lineas, cierre)
    estado = _estado_de(lineas, idx) if idx >= 0 else ""
    escala = _escala_de(lineas, idx if idx >= 0 else None, espanol=True)
    consolidado = _ambito_de_pagina(lineas[:idx]) if idx >= 0 else None
    for c in columnas:
        if c.fin is not None:
            c.periodo = Periodo.instante(c.fin) if estado == "balance" else Periodo.de_meses(c.fin, meses or 12)
    cuerpo = lineas[idx + 1:] if idx >= 0 else []
    if idx < 0 and previa is not None and not any(_TITULO_ES.search(l.texto) for l in lineas[:3]) \
            and _alinea(lineas, previa.columnas, previa.anclas):
        # la continuación de un estado no repite su cabecera: hereda columnas, periodos, estado y escala
        columnas, anclas, estado = previa.columnas, previa.anclas, previa.estado
        escala = _escala_de(lineas, None, espanol=True, por_defecto=previa.escala)
        cuerpo = lineas
    # formato de la página: español salvo que traiga cifras inequívocamente inglesas («1,914,716.70 €»)
    locale_es = _locale_es(lineas) or not _locale_en(lineas)
    filas: List[Fila] = []
    contexto = ""
    recien = False           # la línea anterior acaba de dar una fila: la siguiente puede ser el resto de su rótulo
    for l in _fundir(_sin_notas(cuerpo, columnas, anclas)):
        numericos = [t for t in l.tokens if _es_numero(t.texto)]
        rotulo = " ".join(t.texto for t in l.tokens if t not in numericos).strip(" :")
        sigue, recien = recien, False
        if not numericos:
            if rotulo and sigue and _CONTINUA.search(filas[-1].rotulo) and filas[-1].rect[1] - l.y1 < max(l.y1 - l.y0, 1):
                f = filas[-1]
                f.rotulo = f"{f.rotulo} {rotulo}"
                for cand in f.celdas.values():
                    cand.rotulo = f.rotulo
                if _ENCABEZADO.match(limpiar_rotulo_es(f.rotulo)):
                    contexto = f.rotulo
            elif rotulo and not _PREFIJO.match(rotulo):
                # una partida sin importe del modelo normalizado («2. Deudas con entidades de crédito») no es un epígrafe
                contexto = rotulo
            continue
        if not rotulo:
            continue
        bloque = contexto + " " + rotulo
        es_por_accion = bool(re.search(r"(?i)por acci[óo]n", bloque)) and \
            not re.search(r"(?i)n[úu]mero (?:medio )?(?:ponderado )?de acciones|acciones (?:medias|en circulaci)", bloque)
        celdas: Dict[int, Candidato] = {}
        for t in numericos:
            v, guion, pct = _valor_de(t.texto, locale_es)
            if v is None:
                continue
            j = _destino(columnas, anclas, t.xc)
            if j is None:
                continue
            col = columnas[j]
            celdas[j] = Candidato(documento=documento, pagina=numero, rotulo=rotulo, contexto=contexto,
                                  periodo=col.periodo, etiqueta_columna=col.etiqueta,
                                  valor=v if (es_por_accion or pct) else v * escala, crudo=t.texto, escala=escala,
                                  por_accion=es_por_accion, guion=guion, porcentaje=pct,
                                  rect=(t.x0, t.y0, t.x1, t.y1), rect_fila=(l.x0, l.y0, l.x1, l.y1), moneda=moneda)
        if celdas:
            filas.append(Fila(rotulo=rotulo, contexto=contexto, celdas=celdas, rect=(l.x0, l.y0, l.x1, l.y1)))
            recien = True
        if _ENCABEZADO.match(limpiar_rotulo_es(rotulo)):
            contexto = rotulo
    return PaginaLeida(numero=numero, lineas=lineas, escala=escala, titulo=titulo, columnas=columnas, filas=filas,
                       ancho=ancho, alto=alto, estado=estado, consolidado=consolidado, anclas=list(anclas))


# ---------------------------------------------------------------------------
# Hojas de cálculo
# ---------------------------------------------------------------------------

def extraer_xlsx(adjunto: Adjunto) -> List[PaginaLeida]:
    import openpyxl
    wb = openpyxl.load_workbook(str(adjunto.ruta), data_only=True)
    salida = []
    for numero, ws in enumerate(wb.worksheets, 1):
        filas_crudas = list(ws.iter_rows(values_only=True))
        texto = "\n".join(" ".join("" if c is None else str(c) for c in f) for f in filas_crudas[:12])
        escala = 1_000 if "in thousands" in texto.lower() else 1_000_000 if "in millions" in texto.lower() else 1
        # cabecera: fila de tipos (opcional), fila de meses, fila de años
        idx_anios = None
        for i, f in enumerate(filas_crudas[:12]):
            if sum(isinstance(c, (int, float)) and 1990 <= c <= 2100 for c in f) >= 2:
                idx_anios = i; break
        if idx_anios is None:
            continue
        fila_anios = filas_crudas[idx_anios]
        fila_meses = filas_crudas[idx_anios - 1] if idx_anios >= 1 else ()
        fila_tipos = filas_crudas[idx_anios - 2] if idx_anios >= 2 else ()
        tipo_actual = ""
        columnas: List[Columna] = []
        indices_col: List[int] = []
        for j, c in enumerate(fila_anios):
            if j < len(fila_tipos) and fila_tipos[j]:
                tipo_actual = str(fila_tipos[j])
            if not (isinstance(c, (int, float)) and 1990 <= c <= 2100):
                continue
            mes_txt = str(fila_meses[j]).strip(" ,") if j < len(fila_meses) and fila_meses[j] else ""
            mm = re.fullmatch(r"(" + _MES + r")\s+(\d{1,2})", mes_txt)
            mes = MESES[mm.group(1)] if mm else 12
            fin = _fin_mes(int(c), mes)
            t = tipo_actual.lower()
            if "three" in t: p = Periodo.de_meses(fin, 3)
            elif "six" in t: p = Periodo.de_meses(fin, 6)
            elif "nine" in t: p = Periodo.de_meses(fin, 9)
            elif "twelve" in t or "year" in t: p = Periodo.de_meses(fin, 12)
            elif re.search(r"(?i)balance", ws.title): p = Periodo.instante(fin)
            else: p = None
            columnas.append(Columna(periodo=p, xc=j, etiqueta=f"{tipo_actual} {mes_txt}, {int(c)}".strip(), fin=fin))
            indices_col.append(j)
        filas: List[Fila] = []
        contexto = ""
        por_accion_hoja = "per share" in texto.lower()
        for i, f in enumerate(filas_crudas[idx_anios + 1:], idx_anios + 2):
            rotulo = next((str(c).strip(" :") for c in f[:4] if isinstance(c, str) and c.strip()), "")
            numericos = {j: f[j] for j in indices_col if j < len(f) and isinstance(f[j], (int, float))}
            if not numericos:
                if rotulo:
                    contexto = rotulo
                continue
            if not rotulo:
                continue
            es_por_accion = _es_por_accion(contexto + " " + rotulo)
            celdas = {}
            for k, j in enumerate(indices_col):
                if j not in numericos:
                    continue
                v = float(numericos[j])
                col = columnas[k]
                ref = f"{ws.title}!{openpyxl.utils.get_column_letter(j + 1)}{i}"
                celdas[k] = Candidato(documento=adjunto.nombre, pagina=numero, rotulo=rotulo, contexto=contexto,
                                      periodo=col.periodo, etiqueta_columna=col.etiqueta,
                                      valor=v if es_por_accion else v * escala, crudo=str(numericos[j]), escala=escala,
                                      por_accion=es_por_accion, guion=False, porcentaje=False,
                                      rect=(i, j, i, j), rect_fila=(i, 0, i, len(f)), referencia=ref)
            filas.append(Fila(rotulo=rotulo, contexto=contexto, celdas=celdas, rect=(i, 0, i, len(f))))
        salida.append(PaginaLeida(numero=numero, lineas=[], escala=escala, titulo=ws.title, columnas=columnas, filas=filas))
    return salida
