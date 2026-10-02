"""El expediente: qué ha adjuntado el analista, qué es cada cosa y en qué orden va.

El sistema no se fía del nombre del fichero —EDGAR sirve los PDF con un UUID— ni
de la fecha de creación del PDF, que es la de la descarga. La fecha que importa
es la que el propio documento declara de sí mismo: el periodo que cubre y el día
en que se firmó o se publicó. Las dos se leen de sus páginas, y con ellas el
expediente queda en orden cronológico sin que nadie lo ordene a mano.

Cada clasificación es una inferencia y viaja con su certeza y su motivo (regla
3): un 10-K se reconoce por su portada con certeza alta; un fichero de cuentas
sin portada, por su primera tabla, con certeza media; y lo que no se reconoce
queda como «desconocido» a la vista, no se cuela en un apartado por descarte.

Cuando el emisor presenta ante la SEC, cada 10-K, 10-Q y DEF 14A del expediente
se cruza con los depósitos del emisor por su número de acceso —que EDGAR deja
en el título interno del PDF— y se comprueba que no haya uno más reciente
depositado que no esté en el expediente. Un informe hecho con el 10-Q anterior
al último es un informe desactualizado, y eso hay que decirlo antes de emitir.
"""

from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass, field
from datetime import date
from enum import Enum
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Tuple

from .hechos import Certeza

__all__ = ["Adjunto", "Apartado", "Aviso", "Expediente", "Tipo", "cargar", "clasificar"]


class Tipo(str, Enum):
    K10 = "10-K"
    Q10 = "10-Q"
    DEF14A = "DEF 14A"
    CARTA = "Carta a accionistas"          # anexo 99.1 del 8-K de resultados
    CALL = "Transcripción de la call"       # documento de un tercero, no depositado
    FINWEB = "Cuentas trimestrales (web del emisor)"
    NOTA = "Nota de resultados (anexo 99.1 del 8-K)"
    TABLAS = "Cuentas del anexo 99.1 (8-K)"
    PRESENTACION = "Presentación de resultados"
    OCHOK = "Hecho relevante (8-K)"         # la carátula del 8-K: el contenido va en sus anexos
    XLSX = "Hoja de cálculo de cuentas"
    # emisores españoles (BME): las cuentas anuales auditadas y el informe financiero semestral, en PDF
    CCAA = "Cuentas anuales"
    SEMESTRAL = "Informe financiero semestral"
    # otras comunicaciones de BME que el sistema trae y usa sin leer cifras de ellas
    PARTICIPACIONES = "Participaciones significativas"
    INCORPORACION = "Documento de incorporación"
    # «otra información relevante» o «información privilegiada»: citable por el analista, nunca fuente de cifras
    COMUNICACION = "Comunicación al mercado"
    DESCONOCIDO = "desconocido"


class Apartado(str, Enum):
    ANUALES = "anuales"
    TRIMESTRALES = "trimestrales"
    GUIDANCE = "guidance"


APARTADO_DE = {
    Tipo.K10: Apartado.ANUALES, Tipo.DEF14A: Apartado.ANUALES,
    Tipo.Q10: Apartado.TRIMESTRALES, Tipo.FINWEB: Apartado.TRIMESTRALES, Tipo.XLSX: Apartado.TRIMESTRALES,
    Tipo.CARTA: Apartado.GUIDANCE, Tipo.CALL: Apartado.GUIDANCE,
    Tipo.NOTA: Apartado.GUIDANCE, Tipo.PRESENTACION: Apartado.GUIDANCE, Tipo.TABLAS: Apartado.TRIMESTRALES,
    Tipo.CCAA: Apartado.ANUALES, Tipo.SEMESTRAL: Apartado.TRIMESTRALES, Tipo.COMUNICACION: Apartado.GUIDANCE,
}

# la clave con la que cada tipo se nombra en el expediente y en la lista de documentos del paso 1 (`documentos.py`)
CLAVE_DE = {Tipo.K10: "10K", Tipo.Q10: "10Q", Tipo.DEF14A: "PROXY", Tipo.CARTA: "CARTA",
            Tipo.CALL: "CALL", Tipo.FINWEB: "FINWEB", Tipo.XLSX: "XLSX", Tipo.NOTA: "NOTA",
            Tipo.TABLAS: "TABLAS", Tipo.PRESENTACION: "SLIDES", Tipo.CCAA: "CCAA", Tipo.SEMESTRAL: "SEMESTRAL",
            Tipo.PARTICIPACIONES: "PARTICIPACIONES", Tipo.INCORPORACION: "INCORPORACION", Tipo.COMUNICACION: "COMUNICACION"}
TIPO_DE_CLAVE = {v: k for k, v in CLAVE_DE.items()}

MESES = {m: i for i, m in enumerate(
    ["January", "February", "March", "April", "May", "June", "July",
     "August", "September", "October", "November", "December"], 1)}
_FECHA_LARGA = re.compile(r"(January|February|March|April|May|June|July|August|September|"
                          r"October|November|December)\s+(\d{1,2}),?\s+(\d{4})", re.I)   # las transcripciones van en versales
MESES_ES = {m: i for i, m in enumerate(["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto",
                                        "septiembre", "octubre", "noviembre", "diciembre"], 1)}
MESES_ES["setiembre"] = 9
# «31 de diciembre de 2025»; la capa de texto a veces parte el año («de 202 5») y el día («3 1 de»)
_FECHA_ES = "(?:\\d\\s?)?\\d\\s+de\\s+(?:" + "|".join(MESES_ES) + ")\\s+(?:de|del)\\s+(?:\\d\\s?){3}\\d\\b"


def _fecha_larga(texto: str) -> Optional[date]:
    m = _FECHA_LARGA.search(texto)
    if not m:
        return None
    return date(int(m.group(3)), MESES[m.group(1).capitalize()], int(m.group(2)))


def _fecha_es(texto: str) -> Optional[date]:
    """Una fecha en español: larga («31 de diciembre de 2025») o numérica («30/06/2025», «30.06.2025»)."""
    m = re.search(r"((?:\d\s?)?\d)\s+de\s+(" + "|".join(MESES_ES) + r")\s+(?:de|del)\s+((?:\d\s?){3}\d)\b", texto or "", re.I)
    try:
        if m:
            return date(int(m.group(3).replace(" ", "")), MESES_ES[m.group(2).lower()], int(m.group(1).replace(" ", "")))
        m = re.search(r"\b(\d{1,2})[/.](\d{1,2})[/.]((?:19|20)\d{2})\b", texto or "")
        if m:
            return date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
    except ValueError:
        return None
    return None


def _fin_trimestre(trimestre: int, anio: int, cierre_mes: int = 12) -> date:
    """Fin del trimestre fiscal `trimestre` del ejercicio que cierra en `cierre_mes`."""
    import calendar
    mes = (cierre_mes - 12 + 3 * trimestre - 1) % 12 + 1
    a = anio if mes <= cierre_mes or cierre_mes == 12 else anio - 1
    return date(a, mes, calendar.monthrange(a, mes)[1])


@dataclass
class Adjunto:
    ruta: Path
    huella: str
    paginas: List[str]                       # texto por página (PDF) o por hoja (XLSX)
    tipo: Tipo
    apartado: Optional[Apartado]
    certeza: Certeza
    motivo: str
    periodo_fin: Optional[date] = None       # cierre del periodo que cubre
    fecha: Optional[date] = None             # fecha del documento: firma, publicación o junta
    accession: str = ""                      # número de acceso EDGAR, si el PDF lo trae
    hojas: List[str] = field(default_factory=list)   # solo XLSX
    verificado_en_edgar: Optional[bool] = None       # None = no aplica o no se comprobó
    tambien: Tuple["Tipo", ...] = ()                 # otros documentos del catálogo que trae dentro (`tambien()`)

    @property
    def nombre(self) -> str:
        return self.ruta.name

    @property
    def clave(self) -> str:
        base = CLAVE_DE.get(self.tipo, "DOC")
        if self.periodo_fin:
            base += "_" + self.periodo_fin.strftime("%Y%m%d")
        # con la huella: tipo y periodo no bastan para distinguir dos documentos. Dos notas de resultados sin periodo
        # declarado daban la misma clave y la segunda pisaba a la primera en las páginas leídas del contraste
        return base + ("_" + self.huella[:8] if self.huella else "")

    @property
    def orden(self) -> Tuple[date, date, int]:
        """Cronológico: por cierre del periodo, luego por fecha del documento; los 10-K/10-Q al final de su periodo."""
        pf = self.periodo_fin or self.fecha or date.min
        f = self.fecha or pf
        peso = 1 if self.tipo in (Tipo.K10, Tipo.Q10, Tipo.DEF14A, Tipo.CCAA, Tipo.SEMESTRAL) else 0
        return (pf, f, peso)


@dataclass
class Aviso:
    gravedad: str        # «aviso» | «grave»
    texto: str


@dataclass
class Expediente:
    ticker: str
    adjuntos: List[Adjunto]
    avisos: List[Aviso] = field(default_factory=list)

    def de_tipo(self, tipo: Tipo) -> List[Adjunto]:
        return [a for a in self.adjuntos if a.tipo == tipo]

    def de_apartado(self, apartado: Apartado) -> List[Adjunto]:
        return [a for a in self.adjuntos if a.apartado == apartado]

    def faltan(self) -> List[str]:
        return [ap.value for ap in (Apartado.ANUALES, Apartado.TRIMESTRALES, Apartado.GUIDANCE)
                if not self.de_apartado(ap)]

    def ultimo_periodo(self) -> Optional[date]:
        fines = [a.periodo_fin for a in self.adjuntos if a.periodo_fin]
        return max(fines) if fines else None


# ---------------------------------------------------------------------------
# Clasificación
# ---------------------------------------------------------------------------

def clasificar(p1: str, titulo_pdf: str, primeras: str, n_paginas: Optional[int] = None) -> dict:
    """Qué documento es, por lo que dice de sí mismo. Devuelve tipo, periodo, fecha, certeza, motivo. `n_paginas`: el
    largo del documento, lo único que separa la carta con que un emisor español remite sus cuentas de una comunicación."""
    accession = titulo_pdf.strip() if re.fullmatch(r"\d{10}-\d{2}-\d{6}", titulo_pdf.strip()) else ""
    if "FORM 10-K" in p1:
        m = re.search(r"fiscal year ended\s+(\w+ \d{1,2}, \d{4})", p1)
        return dict(tipo=Tipo.K10, periodo_fin=_fecha_larga(m.group(1)) if m else None, accession=accession,
                    certeza=Certeza.ALTA, motivo="La portada dice «FORM 10-K» y declara el ejercicio fiscal cerrado.")
    if "FORM 10-Q" in p1:
        m = re.search(r"quarterly period ended\s+(\w+ \d{1,2}, \d{4})", p1)
        return dict(tipo=Tipo.Q10, periodo_fin=_fecha_larga(m.group(1)) if m else None, accession=accession,
                    certeza=Certeza.ALTA, motivo="La portada dice «FORM 10-Q» y declara el trimestre cerrado.")
    # la portada que imprime EDGAR: la carátula del formulario, antes de la maqueta del emisor
    if "SCHEDULE 14A" in p1 and "Proxy Statement" in p1:
        m = re.search(r"(?:to be held|will be held)(?: on)?\s+(?:\w+,\s+)?(\w+ \d{1,2}, \d{4})", primeras, re.I)
        return dict(tipo=Tipo.DEF14A, periodo_fin=None, fecha=_fecha_larga(m.group(1)) if m else None,
                    accession=accession, certeza=Certeza.ALTA,
                    motivo="La carátula del depósito dice «SCHEDULE 14A · Proxy Statement»: accionariado, consejo y retribución.")
    if "Proxy" in p1 and "Annual" in p1:
        m = re.search(r"(?:to be held|will be held)(?: on)?\s+(?:\w+,\s+)?(\w+ \d{1,2}, \d{4})", primeras, re.I)
        return dict(tipo=Tipo.DEF14A, periodo_fin=None, fecha=_fecha_larga(m.group(1)) if m else None,
                    accession=accession, certeza=Certeza.ALTA,
                    motivo="La portada dice «Proxy Statement and Notice of Annual Meeting»: accionariado, consejo y retribución.")
    # la proxy que maqueta el emisor abre con una portada de imagen, sin texto: lo que la nombra es la carta de
    # convocatoria de las páginas siguientes. Sin esto, un documento de 120 páginas se queda en «desconocido».
    if (re.search(r"NOTICE OF (?:THE )?(?:\d{4} )?ANNUAL MEETING OF (?:STOCK|SHARE)HOLDERS", primeras, re.I)
            and re.search(r"PROXY STATEMENT", primeras, re.I)):
        m = re.search(r"(?:to be held|will be held)(?: on)?\s+(?:\w+,\s+)?(\w+ \d{1,2}, \d{4})", primeras, re.I)
        return dict(tipo=Tipo.DEF14A, periodo_fin=None, fecha=_fecha_larga(m.group(1)) if m else None,
                    accession=accession, certeza=Certeza.ALTA,
                    motivo="Sus primeras páginas traen el aviso de convocatoria de la junta y el «Proxy Statement»: "
                           "es la proxy maquetada por el emisor, cuya portada es una imagen sin texto.")
    # la carátula del 8-K: el contenido va en los anexos, pero reconocerla evita que un depósito quede «desconocido»
    if re.search(r"\bFORM 8-K\b", p1) and re.search(r"CURRENT REPORT", p1, re.I):
        m = re.search(r"(\w+ \d{1,2}, \d{4})\s+Date of Report", " ".join(p1.split()))
        return dict(tipo=Tipo.OCHOK, periodo_fin=None, fecha=_fecha_larga(m.group(1)) if m else _fecha_larga(p1),
                    accession=accession, certeza=Certeza.ALTA,
                    motivo="La carátula dice «FORM 8-K · CURRENT REPORT»: es el hecho relevante. Las cifras van en sus "
                           "anexos (99.1), no en la carátula.")
    es = _clasificar_es(p1, primeras, n_paginas)
    if es is not None:
        es["accession"] = accession
        return es
    if "Fellow shareholders" in p1:
        m = re.search(r"\bQ([1-4])\b(?:'|’)?(\d{2})?\s+revenue", p1)
        fecha = _fecha_larga(p1)
        pf = None
        if m and fecha:
            anio = 2000 + int(m.group(2)) if m.group(2) else fecha.year
            pf = _fin_trimestre(int(m.group(1)), anio)
        return dict(tipo=Tipo.CARTA, periodo_fin=pf, fecha=fecha, certeza=Certeza.ALTA,
                    motivo="Empieza «Fellow shareholders» con la tabla de resultados y previsión: es el anexo 99.1 del 8-K de resultados.")
    # dos proveedores, dos portadas: «Earnings Call Transcripts» (S&P Global) y «EDITED TRANSCRIPT … Earnings Call» (LSEG)
    if re.search(r"Earnings Call Transcript", p1, re.I) or (re.search(r"\bTRANSCRIPT\b", p1, re.I) and re.search(r"Earnings Call", p1, re.I)):
        m = re.search(r"FQ([1-4]) (\d{4}) Earnings Call", p1)
        fecha = _fecha_larga(p1)
        # solo se fecha el periodo cuando la portada lo dice con el trimestre fiscal; el trimestre natural de otras
        # portadas no equivale al fiscal —Qualcomm cierra en septiembre—, y un cierre inventado es peor que ninguno
        pf = _fin_trimestre(int(m.group(1)), int(m.group(2))) if m else None
        return dict(tipo=Tipo.CALL, periodo_fin=pf, fecha=fecha, certeza=Certeza.ALTA,
                    motivo="La portada la titula transcripción de la «Earnings Call»"
                           + (f", del trimestre fiscal {m.group(1)} de {m.group(2)}" if m else "")
                           + ". Documento de un tercero (S&P Global, LSEG), no depositado en la SEC.")
    if "Consolidated Statements of Operations" in p1 and "unaudited" in p1:
        fechas = [_fecha_larga(f"{m} {a}") for m, a in _pares_mes_anio(p1)]
        pf = max(f for f in fechas if f) if any(fechas) else None
        return dict(tipo=Tipo.FINWEB, periodo_fin=pf, certeza=Certeza.MEDIA,
                    motivo="Sin portada: empieza por la cuenta de resultados trimestral «(unaudited)», como el fichero de cuentas de la web de relación con inversores.")
    # los estados condensados del anexo 99.1 del 8-K: en versales y sin «unaudited» en la portada, no casan con el fichero de la web
    if re.search(r"CONDENSED CONSOLIDATED STATEMENTS? OF OPERATIONS", p1) and re.search(r"FINANCIAL RESULTS|UNAUDITED", p1):
        m = re.search(r"[Ee]nded\s+(\w+ \d{1,2}, \d{4})", re.sub(r"\s+", " ", primeras))   # el cierre puede venir partido entre líneas
        return dict(tipo=Tipo.TABLAS, periodo_fin=_fecha_larga(m.group(1)) if m else None, accession=accession, certeza=Certeza.ALTA,
                    motivo="Empieza por los estados condensados del trimestre en versales: son las tablas que acompañan a la nota de resultados (anexo 99.1 del 8-K)"
                           + (f", del periodo cerrado el {m.group(1)}." if m else "; la portada no declara el cierre del periodo."))
    # la nota de prensa de resultados: titular «… Announces … Results» y el bloque de relación con inversores
    if re.search(r"Announces[^.\n]{0,80}Results", p1) and re.search(r"Investor Relations|Earnings Release", p1, re.I):
        return dict(tipo=Tipo.NOTA, periodo_fin=None, fecha=_fecha_larga(p1), accession=accession, certeza=Certeza.ALTA,
                    motivo="Titular «Announces … Results» con el contacto de relación con inversores: es la nota de resultados (anexo 99.1 del 8-K). "
                           "El cierre del periodo no se declara literalmente en la portada, así que no se le atribuye ninguno.")
    # la misma nota con los otros titulares habituales («reports / posts / delivers / releases … results», en cualquier
    # caja), con el bloque de relación con inversores o con la fórmula «today announced … results» como segunda prueba.
    # Sin fecha de la portada: la primera que trae suele ser el cierre del trimestre («quarter ended …») y las
    # siguientes, dividendo y call; la del documento la pone su depósito
    if (re.search(r"\b(?:reports|posts|delivers|releases|announces)\b[^.\n]{0,80}\bresults\b", p1, re.I)
            and (re.search(r"Investor Relations|Earnings Release", p1, re.I)
                 or re.search(r"\btoday (?:announced|reported)\b[^.\n]{0,80}\bresults\b", p1, re.I))):
        return dict(tipo=Tipo.NOTA, periodo_fin=None, fecha=None, accession=accession, certeza=Certeza.ALTA,
                    motivo="Titular de resultados («reports / posts / delivers / releases … results») con el bloque de relación con inversores "
                           "o el «today announced … results»: es la nota de resultados (anexo 99.1 del 8-K). "
                           "Ni el cierre del periodo ni la fecha de publicación se declaran solos en la portada, así que no se le atribuyen: "
                           "la fecha la pone su depósito de EDGAR.")
    # la presentación de resultados: el rótulo de la primera diapositiva —que puede venir partido en dos líneas—,
    # o el título que el propio PDF declara. No vale buscarlo en toda la página: debajo va el descargo legal, que en
    # una portada real ocupa más que el rótulo. Lo que distingue a la diapositiva es que su rótulo es una línea corta.
    lineas = [l.strip() for l in p1.splitlines() if l.strip()]
    rotulo = " ".join(lineas[:2])
    por_titulo = re.search(r"Earnings Presentation", titulo_pdf, re.I)
    if (len(rotulo) < 90 and re.search(r"(?:Q[1-4]|First|Second|Third|Fourth)[^\n]{0,40}\bEarnings\b", rotulo)) or por_titulo:
        return dict(tipo=Tipo.PRESENTACION, periodo_fin=None, fecha=_fecha_larga(p1), certeza=Certeza.MEDIA,
                    motivo=("El PDF se titula «" + titulo_pdf.strip() + "»" if por_titulo
                            else "La primera diapositiva se rotula con el trimestre y «Earnings»")
                           + ": es la presentación de resultados. Material de apoyo: las cifras se toman de las cuentas.")
    return dict(tipo=Tipo.DESCONOCIDO, accession=accession, certeza=Certeza.BAJA,
                motivo="Ninguna pista reconocible en la primera página.")


# Cuentas españolas. En la primera página vale la mención débil («primer semestre», «cuentas anuales»): es la carta
# con la que el emisor las remite a BME o la portada del informe. En las siguientes solo la fuerte —el informe de
# auditoría o de revisión, los estados intermedios—, porque una presentación de resultados anuales habla también de su
# primer semestre y una semestral cita las cuentas anuales del año anterior. Si hay dos, manda la que aparece antes.
_ES_SEMESTRAL_FUERTE = r"estados financieros intermedios|informe financiero semestral|periodo de seis meses (?:terminado|cerrado|finalizado)|informe de revisi[óo]n limitada"
_ES_CCAA_FUERTE = (r"informe de auditor[íi]a (?:independiente )?de (?:las )?cuentas anuales"
                   r"|cuentas anuales(?: consolidadas)?(?: e informe de gesti[óo]n)? (?:al|a|correspondientes al ejercicio|del ejercicio)\b")
_ES_SEMESTRAL = _ES_SEMESTRAL_FUERTE + r"|primer semestre|1er semestre|informaci[óo]n financiera (?:semestral|intermedia)"
_ES_CCAA = _ES_CCAA_FUERTE + r"|cuentas anuales|informaci[óo]n financiera (?:anual|del ejercicio)"
_CIERRE_CCAA = (r"ejercicio (?:anual )?(?:terminado|cerrado|finalizado) (?:el|a|al|en) (" + _FECHA_ES + ")",
                r"cuentas anuales(?: consolidadas)?(?: e informe de gesti[óo]n)? (?:al|a) (" + _FECHA_ES + ")",
                r"balance[^.]{0,60}?\b(?:a|al) (" + _FECHA_ES + ")")
_CIERRE_SEMESTRAL = (r"seis meses (?:terminado|cerrado|finalizado) (?:el|a|al|en) (" + _FECHA_ES + ")",
                     r"(?:balance|situaci[óo]n financiera)[^.]{0,60}?\b(?:a|al) (" + _FECHA_ES + ")",
                     r"estados financieros intermedios[^.]{0,80}?\b(?:a|al) (" + _FECHA_ES + ")")


_ES_COMUNICACION = r"otra informaci[óo]n relevante|informaci[óo]n privilegiada"


def _clasificar_es(p1: str, primeras: str, n_paginas: Optional[int] = None) -> Optional[dict]:
    """Las cuentas anuales, el semestral o una comunicación de un emisor español, por lo que dicen de sí mismos; None si
    no son ninguna. Cuentas y semestral llegan a BME con la misma carta que una comunicación: lo que las separa es el
    largo (unos estados financieros no caben en `umbrales.comunicacion_bme_paginas_max` páginas)."""
    t1, tp = " ".join(p1.split()), " ".join(primeras.split())
    if n_paginas is not None and re.search(_ES_COMUNICACION, t1[:300], re.I):
        from ..umbrales import umbral
        if n_paginas <= int(umbral("comunicacion_bme_paginas_max")):
            return dict(tipo=Tipo.COMUNICACION, periodo_fin=None, fecha=_fecha_es(t1[:200]), certeza=Certeza.ALTA,
                        motivo=f"La primera página es una comunicación al mercado («{re.search(_ES_COMUNICACION, t1, re.I).group(0)}») "
                               f"de {n_paginas} página{'s' if n_paginas != 1 else ''}: se cita, pero no se toman cifras de ella.")
    tipo = None
    for texto, semestral, anual in ((t1, _ES_SEMESTRAL, _ES_CCAA), (tp, _ES_SEMESTRAL_FUERTE, _ES_CCAA_FUERTE)):
        s, a = re.search(semestral, texto, re.I), re.search(anual, texto, re.I)
        if s or a:
            m = s if (s and (not a or s.start() <= a.start())) else a
            tipo, frase = (Tipo.SEMESTRAL if m is s else Tipo.CCAA), m.group(0)
            break
    if tipo is None:
        return None
    fin, deducido = None, False
    for patron in (_CIERRE_SEMESTRAL if tipo is Tipo.SEMESTRAL else _CIERRE_CCAA):
        m = re.search(patron, t1, re.I) or re.search(patron, tp, re.I)
        if m:
            fin = _fecha_es(m.group(1))
            break
    if fin is None and tipo is Tipo.SEMESTRAL:
        # «primer semestre de 2025» sin fecha escrita: el cierre es el 30 de junio, y se dice que es una deducción
        m = re.search(r"(?:primer|1er) semestre (?:del ejercicio |de |del )?((?:19|20)\d{2})", tp, re.I)
        if m:
            fin, deducido = date(int(m.group(1)), 6, 30), True
    # la carta con que el emisor lo remite a BME lleva la fecha de la comunicación en su primera línea
    fecha = _fecha_es(t1[:200]) if re.search(r"art[íi]culo 17 del Reglamento|informaci[óo]n relevante|informaci[óo]n privilegiada",
                                              t1, re.I) else None
    nombre = "las cuentas anuales" if tipo is Tipo.CCAA else "el informe financiero semestral"
    cierre = (f", con cierre el {fin:%d/%m/%Y}" + (" (deducido de «primer semestre»: no escribe la fecha)" if deducido else "")
              if fin else "; no declara la fecha de cierre, así que no se le atribuye ninguna")
    return dict(tipo=tipo, periodo_fin=fin, fecha=fecha, certeza=Certeza.MEDIA if deducido else Certeza.ALTA,
                motivo=f"El documento dice ser {nombre} («{frase}»){cierre}.")


# el rótulo de un estado de resultados como encabezado de su propia línea: mencionarlo en la prosa no es traerlo
_ESTADO_DE_RESULTADOS = re.compile(r"^\s*(?:CONDENSED\s+)?CONSOLIDATED\s+STATEMENTS?\s+OF\s+(?:OPERATIONS|INCOME|EARNINGS)\b", re.I | re.M)


def tambien(tipo: Tipo, paginas: Sequence[str]) -> Tuple[Tipo, ...]:
    """Los otros documentos del catálogo que un adjunto trae dentro, además del que dice ser su portada.

    La nota de resultados y la carta a accionistas suelen llevar detrás los estados condensados del trimestre: son las
    mismas cuentas que el anexo 99.1 publica sueltas en otros emisores, el contraste ya las lee de sus páginas, y la
    lista de documentos no puede decir que faltan. La portada se salta: ya dice lo que el documento es.
    """
    if tipo in (Tipo.NOTA, Tipo.CARTA) and any(_ESTADO_DE_RESULTADOS.search(p) for p in paginas[1:]):
        return (Tipo.TABLAS,)
    return ()


def _pares_mes_anio(texto: str) -> List[Tuple[str, str]]:
    """Cabeceras de dos líneas: «March 31, June 30, …» y debajo «2025 2025 …», emparejadas por posición."""
    lineas = texto.splitlines()
    for i, l in enumerate(lineas[:-1]):
        meses = re.findall(r"((?:January|February|March|April|May|June|July|August|September|October|November|December) \d{1,2}),", l)
        anios = re.findall(r"\b(\d{4})\b", lineas[i + 1])
        if meses and len(meses) == len(anios):
            return list(zip(meses, anios))
    return []


# ---------------------------------------------------------------------------
# Carga
# ---------------------------------------------------------------------------

def _leer_pdf(ruta: Path) -> Tuple[List[str], str]:
    import pypdfium2 as pdfium
    doc = pdfium.PdfDocument(str(ruta))
    try:
        paginas = []
        for i in range(len(doc)):
            pagina = doc[i]
            texto = pagina.get_textpage()
            paginas.append(texto.get_text_range())
            texto.close()
            pagina.close()
        titulo = doc.get_metadata_value("Title") or ""
    finally:
        # sin cerrarlo, pdfium deja el fichero abierto mientras viva el proceso: en Windows el analista no puede
        # quitar ni sustituir un documento que el servidor ya haya leído, y la memoria del servidor no baja nunca
        doc.close()
    return paginas, titulo


def _leer_xlsx(ruta: Path) -> Tuple[List[str], List[str]]:
    import openpyxl
    wb = openpyxl.load_workbook(str(ruta), data_only=True, read_only=True)
    paginas, hojas = [], []
    for ws in wb.worksheets:
        filas = []
        for fila in ws.iter_rows(values_only=True):
            filas.append("\t".join("" if c is None else str(c) for c in fila))
        paginas.append("\n".join(filas)); hojas.append(ws.title)
    return paginas, hojas


def _leer_docx(ruta: Path) -> List[str]:
    """El texto de un .docx con la biblioteca estándar (zip + XML): párrafos separados por líneas, saltos de página como páginas."""
    import zipfile
    from xml.etree import ElementTree as ET
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    with zipfile.ZipFile(str(ruta)) as z:
        raiz = ET.fromstring(z.read("word/document.xml"))
    paginas, actual = [], []
    for parrafo in raiz.iter(f"{{{ns['w']}}}p"):
        if parrafo.find(".//w:br[@w:type='page']", ns) is not None or parrafo.find(".//w:lastRenderedPageBreak", ns) is not None:
            paginas.append("\n".join(actual)); actual = []
        actual.append("".join(t.text or "" for t in parrafo.iter(f"{{{ns['w']}}}t")))
    paginas.append("\n".join(actual))
    return [p for p in paginas if p.strip()] or [""]


def cargar_adjunto(ruta: Path) -> Adjunto:
    ruta = Path(ruta)
    huella = hashlib.sha256(ruta.read_bytes()).hexdigest()
    if ruta.suffix.lower() in (".xlsx", ".xlsm"):
        try:
            paginas, hojas = _leer_xlsx(ruta)
        except Exception as e:  # un fichero vacío o corrupto no es un adjunto: se declara, no se cae
            return Adjunto(ruta=ruta, huella=huella, paginas=[], tipo=Tipo.DESCONOCIDO, apartado=None,
                           certeza=Certeza.BAJA, motivo=f"No se pudo abrir como hoja de cálculo ({e.__class__.__name__}: {e}).")
        fechas = [_fecha_larga(f"{m} {a}") for m, a in _pares_mes_anio(paginas[0].replace("\t", " "))] if paginas else []
        # en una hoja de cálculo las cabeceras vienen en celdas separadas: «March 31,» y «2025» en dos filas
        pf = max((f for f in fechas if f), default=None)
        if pf is None and paginas:
            pf = _fin_por_celdas(paginas[0])
        return Adjunto(ruta=ruta, huella=huella, paginas=paginas, tipo=Tipo.XLSX, apartado=Apartado.TRIMESTRALES,
                       certeza=Certeza.MEDIA, periodo_fin=pf, hojas=hojas,
                       motivo=f"Hoja de cálculo con {len(hojas)} hojas ({', '.join(hojas[:4])}…); sin portada, el periodo es la última columna de la primera hoja.")
    if ruta.suffix.lower() == ".docx":
        # Word: se lee el texto para reconocer el documento y ordenarlo, pero sin coordenadas de página no hay extracción
        # con evidencia ni recorte, así que las cifras no salen de él: el analista debe adjuntar el PDF oficial
        try:
            paginas = _leer_docx(ruta)
        except Exception as e:
            return Adjunto(ruta=ruta, huella=huella, paginas=[], tipo=Tipo.DESCONOCIDO, apartado=None,
                           certeza=Certeza.BAJA, motivo=f"No se pudo abrir como Word ({e.__class__.__name__}: {e}).")
        c = clasificar(paginas[0] if paginas else "", "", "\n".join(paginas[:8]))
        return Adjunto(ruta=ruta, huella=huella, paginas=paginas, tipo=Tipo.DESCONOCIDO, apartado=None, certeza=Certeza.BAJA,
                       periodo_fin=c.get("periodo_fin"), fecha=c.get("fecha"),
                       motivo=f"Word reconocido como «{c['tipo'].value}» ({c['motivo']}); solo sirve para ordenar el expediente: "
                              "las cifras y las evidencias se toman del PDF oficial o de la SEC.")
    if ruta.suffix.lower() != ".pdf":
        return Adjunto(ruta=ruta, huella=huella, paginas=[], tipo=Tipo.DESCONOCIDO, apartado=None,
                       certeza=Certeza.BAJA, motivo=f"Formato {ruta.suffix} no admitido: se aceptan PDF con capa de texto, XLSX y Word (solo para ordenar).")
    paginas, titulo = _leer_pdf(ruta)
    if not any(p.strip() for p in paginas[:3]):
        # regla 12: sin capa de texto no hay cifra con su página y su recorte; se pide, no se adivina (ni OCR)
        return Adjunto(ruta=ruta, huella=huella, paginas=paginas, tipo=Tipo.DESCONOCIDO, apartado=None,
                       certeza=Certeza.BAJA, motivo="sin texto: pide las cifras al analista. El PDF no tiene capa de texto "
                                                    "(está escaneado) y no se lee con OCR.")
    c = clasificar(paginas[0], titulo, "\n".join(paginas[:8]), len(paginas))
    fecha = c.get("fecha")
    if fecha is None and c["tipo"] in (Tipo.K10, Tipo.Q10):
        for t in reversed(paginas):
            m = re.search(r"Dated?:\s*(\w+ \d{1,2}, \d{4})", t)
            if m:
                fecha = _fecha_larga(m.group(1)); break
    return Adjunto(ruta=ruta, huella=huella, paginas=paginas, tipo=c["tipo"], apartado=APARTADO_DE.get(c["tipo"]),
                   certeza=c["certeza"], motivo=c["motivo"], periodo_fin=c.get("periodo_fin"), fecha=fecha,
                   accession=c.get("accession", ""))


def _fin_por_celdas(texto_hoja: str) -> Optional[date]:
    lineas = texto_hoja.splitlines()
    for i, l in enumerate(lineas[:-1]):
        celdas_m = [c.strip(" ,") for c in l.split("\t")]
        celdas_a = [c.strip() for c in lineas[i + 1].split("\t")]
        pares = [(m, a) for m, a in zip(celdas_m, celdas_a)
                 if re.fullmatch(r"(?:January|February|March|April|May|June|July|August|September|October|November|December) \d{1,2}", m)
                 and re.fullmatch(r"\d{4}(?:\.0)?", a)]
        if pares:
            return max(_fecha_larga(f"{m} {a[:4]}") for m, a in pares)
    return None


def declaraciones(carpeta: Path) -> Dict[str, str]:
    """Lo que el analista declaró al adjuntar: fichero → clave de la casilla (`declarado.json` de su propia carpeta).

    Es una declaración, no una lectura del documento: manda siempre lo que el documento diga de sí mismo.
    """
    ruta = Path(carpeta) / "declarado.json"
    if not ruta.is_file():
        return {}
    try:
        import json
        datos = json.loads(ruta.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}
    return {str(k): str(v) for k, v in datos.items()} if isinstance(datos, dict) else {}


def origenes(carpeta: Path) -> Dict[str, dict]:
    """De qué depósito de EDGAR salió cada documento que trajo el sistema (`origen.json` de su propia carpeta).

    Esto no es una inferencia: el fichero se imprimió del documento que EDGAR sirvió en esa dirección, y el depósito
    dice su formulario, su número de acceso y sus fechas. Es la fuente más fuerte que puede tener un adjunto.
    """
    ruta = Path(carpeta) / "origen.json"
    if not ruta.is_file():
        return {}
    try:
        import json
        datos = json.loads(ruta.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}
    return {str(k): v for k, v in datos.items() if isinstance(v, dict)} if isinstance(datos, dict) else {}


def _fecha(valor) -> Optional[date]:
    try:
        return date.fromisoformat(str(valor)) if valor else None
    except ValueError:
        return None


def _de_bme(a: Adjunto, dato: dict) -> None:
    """Lo que BME dice de un documento que trajo el sistema: título, fecha de publicación, número de registro y, en la
    información financiera, el ejercicio y el periodo («AN», «1S»). La portada manda si reconoce el documento; si no
    (un PDF de 262 páginas que abre con el informe de gestión), manda la bolsa, que sabe lo que publicó. La fecha de
    cierre la dice el propio documento: se busca la del ejercicio que declara la bolsa, nunca se supone el 31/12."""
    publicado = _fecha(dato.get("publicado"))
    a.fecha = publicado or a.fecha              # la de publicación en la bolsa, no una fecha cualquiera del texto
    periodo, ejercicio = str(dato.get("periodo") or ""), dato.get("ejercicio")
    tipo = Tipo.CCAA if periodo == "AN" else Tipo.SEMESTRAL if periodo.endswith("S") else \
        {"presentacion": Tipo.PRESENTACION, "participaciones": Tipo.PARTICIPACIONES,
         "incorporacion": Tipo.INCORPORACION, "directivos": Tipo.COMUNICACION,
         "comunicacion": Tipo.COMUNICACION}.get(dato.get("clave"))
    # una comunicación es solo el sobre: lo que la bolsa dice que publicó (participaciones, presentación) es más preciso
    if tipo is not None and a.tipo in (Tipo.DESCONOCIDO, Tipo.COMUNICACION, tipo):
        if a.tipo in (Tipo.DESCONOCIDO, Tipo.COMUNICACION):
            a.certeza = Certeza.ALTA if tipo in (Tipo.CCAA, Tipo.SEMESTRAL) else Certeza.MEDIA
        a.tipo, a.apartado = tipo, APARTADO_DE.get(tipo)
    if a.tipo in (Tipo.CCAA, Tipo.SEMESTRAL) and a.periodo_fin is None and ejercicio:
        a.periodo_fin = _cierre_en_texto(a, int(ejercicio), a.tipo is Tipo.SEMESTRAL)
    a.motivo += (f" Traído de BME: «{dato.get('titulo', '')}», publicado el {publicado:%d/%m/%Y}"
                 + (f" (registro {dato['id']})" if dato.get("id") else "") + "." if publicado else "")


def _cierre_en_texto(a: Adjunto, ejercicio: int, semestral: bool) -> Optional[date]:
    """La fecha de cierre más citada del ejercicio (o del semestre) en el propio documento, detrás de «terminado»,
    «cerrado», «al» o «a»; None si el documento no la dice."""
    from collections import Counter
    patron = re.compile(r"(?:terminad[oa]|cerrad[oa]|finalizad[oa]|\bal|\ba)\s+(?:el\s+)?(" + _FECHA_ES + ")", re.I)
    cuenta: Counter = Counter()
    for texto in a.paginas:
        for m in patron.finditer(" ".join(texto.split())):
            dia = _fecha_es(m.group(1))
            if dia is not None and dia.year == ejercicio and (not semestral or dia.month != 12):
                cuenta[dia] += 1
    return cuenta.most_common(1)[0][0] if cuenta else None


def _de_edgar(a: Adjunto, dato: dict, avisos: List[Aviso]) -> None:
    """Le pone al adjunto lo que EDGAR dice de él: formulario, número de acceso y fechas del depósito."""
    formulario = str(dato.get("formulario") or "")
    presentado, periodo = _fecha(dato.get("presentado")), _fecha(dato.get("periodo"))
    tipo = {"10-K": Tipo.K10, "10-Q": Tipo.Q10, "DEF 14A": Tipo.DEF14A}.get(formulario)
    a.accession = str(dato.get("accession") or a.accession)
    if tipo is None:
        # un anexo 99 del 8-K: EDGAR no dice si es la nota, las tablas o la presentación; eso lo dice su portada
        a.fecha = a.fecha or presentado
        a.motivo += f" Traído de EDGAR: anexo del 8-K {a.accession}, presentado el {presentado:%d/%m/%Y}." if presentado else ""
        # si la portada no trae ninguna pista, su procedencia sí la da: `fuentes.edgar` solo trae anexos 99 del 8-K
        # de resultados (epígrafe 2.02), y el anexo 99 de ese 8-K es la nota. Es una inferencia, y se dice
        from ..fuentes.sec import es_anexo_99
        if a.tipo is Tipo.DESCONOCIDO and formulario == "8-K" and es_anexo_99(str(dato.get("url") or "").rsplit("/", 1)[-1]):
            a.tipo, a.apartado, a.certeza = Tipo.NOTA, APARTADO_DE[Tipo.NOTA], Certeza.MEDIA
            a.motivo = ("Su portada no trae ninguna pista reconocible, pero EDGAR lo sirve como anexo 99 del 8-K de resultados "
                        f"{a.accession}" + (f", presentado el {presentado:%d/%m/%Y}" if presentado else "")
                        + ": se toma como la nota de resultados por su procedencia, no por lo que dice el documento.")
        return
    if a.tipo is not tipo and a.tipo is not Tipo.DESCONOCIDO:
        avisos.append(Aviso("grave", f"{a.nombre}: EDGAR lo sirve como «{formulario}» y su portada dice «{a.tipo.value}»; "
                                     f"se usa el formulario del depósito."))
    a.tipo, a.apartado, a.certeza = tipo, APARTADO_DE.get(tipo), Certeza.ALTA
    if tipo is Tipo.DEF14A:
        a.fecha = periodo or presentado
    else:
        a.periodo_fin = a.periodo_fin or periodo
        a.fecha = a.fecha or presentado
    a.motivo = (f"Traído de EDGAR: depósito {formulario} con número de acceso {a.accession}"
                + (f", presentado el {presentado:%d/%m/%Y}" if presentado else "")
                + (f", periodo {periodo:%d/%m/%Y}" if periodo else "") + ".")


def _declarado(a: Adjunto, clave: str, avisos: List[Aviso]) -> None:
    """Cruza la casilla en la que el analista lo adjuntó con lo que el documento dice ser (regla 3)."""
    tipo = TIPO_DE_CLAVE.get(clave.upper())
    if tipo is None or tipo is a.tipo:
        return
    if a.tipo is Tipo.DESCONOCIDO and a.ruta.suffix.lower() == ".pdf" and a.paginas:
        # el documento no trae ninguna pista reconocible: se usa la declaración del analista, dicha como tal
        a.tipo, a.apartado, a.certeza = tipo, APARTADO_DE.get(tipo), Certeza.MEDIA
        a.motivo = (f"El analista lo adjuntó en la casilla «{tipo.value}» y el documento no trae ninguna pista reconocible "
                    f"en su primera página: se usa la declaración del analista, no una lectura del documento.")
        avisos.append(Aviso("aviso", f"{a.nombre}: se toma como «{tipo.value}» porque lo declaró el analista; el documento no lo dice."))
        return
    avisos.append(Aviso("grave", f"{a.nombre}: lo adjuntaste como «{tipo.value}» y el documento dice ser «{a.tipo.value}» "
                                 f"({a.motivo}). Manda el documento: se usa como «{a.tipo.value}»."))


def cargar(ticker: str, rutas: Iterable[Path], depositos: Optional[Sequence] = None) -> Expediente:
    """Carga, clasifica, deduplica por huella y ordena. Con `depositos` (de `sec.depositos`), cruza con EDGAR."""
    vistos: Dict[str, Adjunto] = {}
    avisos: List[Aviso] = []
    declarado: Dict[Path, Dict[str, str]] = {}
    traido: Dict[Path, Dict[str, dict]] = {}
    for ruta in rutas:
        a = cargar_adjunto(Path(ruta))
        carpeta = Path(ruta).parent
        if carpeta not in declarado:
            declarado[carpeta], traido[carpeta] = declaraciones(carpeta), origenes(carpeta)
        origen = traido[carpeta].get(a.nombre)
        clave = declarado[carpeta].get(a.nombre, "")
        if origen and origen.get("fuente") == "bme":
            _de_bme(a, origen)                # traído de la bolsa española: su fecha de publicación y su registro
        elif origen:
            _de_edgar(a, origen, avisos)      # lo que dice el depósito manda sobre la casilla del analista
        elif clave:
            _declarado(a, clave, avisos)
        a.tambien = tambien(a.tipo, a.paginas)       # después de EDGAR y de la casilla: el tipo ya es el definitivo
        if a.huella in vistos:
            avisos.append(Aviso("aviso", f"{a.nombre} es una copia exacta de {vistos[a.huella].nombre}: se usa una sola vez."))
            continue
        vistos[a.huella] = a
        if a.tipo is Tipo.DESCONOCIDO and not (origen and origen.get("fuente") == "bme"):
            # lo traído de BME sabe lo que es aunque su portada no lo diga (participaciones, documento de incorporación)
            avisos.append(Aviso("grave", f"{a.nombre}: {a.motivo}"))
    adjuntos = sorted(vistos.values(), key=lambda a: a.orden)
    exp = Expediente(ticker=ticker.upper(), adjuntos=adjuntos, avisos=avisos)
    for ap in exp.faltan():
        avisos.append(Aviso("grave", f"Falta el apartado «{ap}» del expediente."))
    if depositos is not None:
        _cruzar_con_edgar(exp, depositos)
    return exp


def _cruzar_con_edgar(exp: Expediente, depositos: Sequence) -> None:
    """Cada 10-K/10-Q/DEF 14A del expediente debe existir en EDGAR, y no debe haber uno más nuevo fuera."""
    por_accession = {d.accession: d for d in depositos}
    for a in exp.adjuntos:
        formulario = {Tipo.K10: "10-K", Tipo.Q10: "10-Q", Tipo.DEF14A: "DEF 14A"}.get(a.tipo)
        if not formulario:
            # un anexo del 8-K o la propia carátula: no hay formulario que casar, pero su número de acceso está en
            # EDGAR igual que el de un 10-K, y decir «—» donde hay depósito es esconder la mitad de la procedencia
            d = por_accession.get(a.accession) if a.accession else None
            if d is not None:
                a.verificado_en_edgar = True
                a.fecha = a.fecha or d.presentado
            continue
        d = por_accession.get(a.accession) if a.accession else None
        if d is None:
            # Un PDF maquetado por el emisor (la proxy de su web) no lleva el número de
            # acceso: se casa por la fecha que EDGAR llama `reportDate` —cierre del
            # periodo en 10-K/10-Q, fecha de la junta en una DEF 14A—.
            fecha_clave = a.fecha if a.tipo is Tipo.DEF14A else a.periodo_fin
            candidatos = [x for x in depositos if x.formulario == formulario and x.periodo == fecha_clave]
            d = max(candidatos, key=lambda x: x.presentado) if candidatos else None
            if d is not None:
                a.accession = d.accession
        if d is not None:
            a.verificado_en_edgar = True
            if a.fecha is None:
                a.fecha = d.presentado
        else:
            a.verificado_en_edgar = False
            exp.avisos.append(Aviso("aviso", f"{a.nombre} ({a.tipo.value}): no se ha podido casar con ningún depósito de EDGAR (ni por número de acceso ni por fecha)."))
    for formulario, tipo in (("10-K", Tipo.K10), ("10-Q", Tipo.Q10), ("DEF 14A", Tipo.DEF14A)):
        en_edgar = [d for d in depositos if d.formulario == formulario]
        en_exp = exp.de_tipo(tipo)
        if en_edgar and en_exp:
            ultimo = max(en_edgar, key=lambda d: d.presentado)
            if all(a.accession != ultimo.accession for a in en_exp):
                exp.avisos.append(Aviso("grave", f"Hay un {formulario} más reciente en EDGAR (presentado el {ultimo.presentado:%d/%m/%Y}, periodo {ultimo.periodo}) que no está en el expediente."))
