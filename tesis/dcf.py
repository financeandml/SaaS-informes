"""El modelo de valoración del analista (sección D, apartados 12–20), leído de su hoja de cálculo.

El analista adjunta un libro Excel con su DCF; el sistema no valora nada: lee lo que el
libro dice y lo imprime como lo que es —supuestos del analista, capa S— con la celda de
origen de cada cifra, y **cuadra** las entradas del modelo que afirman ser un hecho
(ingresos del año base, margen, deuda, caja, acciones) contra los hechos contrastados
del informe. Un modelo que parte de una deuda distinta de la del 10-Q lo dice aquí, con
las dos cifras, antes de que el precio objetivo se imprima.

Se lee por rótulos, no por coordenadas fijas: el libro puede cambiar de fila, pero no
de vocabulario. Cada bloque que no se encuentra se declara en `faltan` y su apartado
sale N/A con motivo. Lo que sí se exige del libro:

- una hoja de **supuestos** (rótulo · valor · origen/justificación);
- una hoja por **escenario** (pesimista, base, optimista) con WACC, crecimiento
  terminal, peso, una fila de años (2026E…) y las filas de proyección y de valoración;
- una hoja de **sensibilidad** (matriz WACC × g) y una de **reverse DCF**;
- una hoja **resumen** con la tabla de escenarios, las referencias de precio, la
  valoración por múltiplos y los anclajes del precio objetivo.

Con `data_only=True` se leen los valores calculados por Excel y, en una segunda
pasada, las fórmulas: una celda con fórmula es un derivado del modelo (D dentro de S);
una celda con número tecleado es un supuesto puro. El informe lo distingue.
"""

from __future__ import annotations

import hashlib
import tempfile
import re
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Tuple

from .formato import numero
from .hechos import Contraste, Hecho, Periodo

__all__ = ["Cuadre", "Escenario", "Modelo", "Multiplo", "Supuesto", "cargar", "cuadrar", "fecha_precio_libro"]


@dataclass(frozen=True)
class Celda:
    hoja: str
    ref: str
    formula: Optional[str] = None      # None = número tecleado (supuesto puro)
    formato: Optional[str] = None      # el formato numérico de la celda («0.0%»): decide la escala al imprimir

    @property
    def cita(self) -> str:
        return f"{self.hoja}!{self.ref}"


@dataclass
class Supuesto:
    rotulo: str
    valor: float
    origen: str                        # la justificación que el analista escribió al lado
    celda: Celda


@dataclass
class Escenario:
    nombre: str
    hoja: str
    wacc: Optional[float]
    g: Optional[float]
    peso: Optional[float]
    anios: List[str]
    filas: Dict[str, List[Optional[float]]]          # rótulo → valores por año, en el orden de `anios`
    resumen: Dict[str, float]                        # rótulo → valor (suma de VP, valor terminal, EV, valor por acción…)
    celdas: Dict[str, Celda] = field(default_factory=dict)
    citas: Dict[str, List[str]] = field(default_factory=dict)   # rótulo → cita de cada columna de `anios`

    def fila(self, patron: str) -> Optional[Tuple[str, List[Optional[float]]]]:
        rx = re.compile(patron, re.I)
        return next(((k, v) for k, v in self.filas.items() if rx.search(k)), None)

    def dato(self, patron: str) -> Optional[Tuple[str, float]]:
        rx = re.compile(patron, re.I)
        return next(((k, v) for k, v in self.resumen.items() if rx.search(k)), None)


@dataclass
class Multiplo:
    rotulo: str
    base: Optional[float]
    actual: Optional[float]
    objetivo: Optional[float]
    valor_accion: Optional[float]
    origen: str
    celda: Celda
    celdas: List[Celda] = field(default_factory=list)   # base, actual, objetivo, valor por acción


@dataclass
class Cuadre:
    rotulo_modelo: str
    valor_modelo: float
    campo: str
    periodo: str
    valor_informe: Optional[float]
    contraste: Contraste
    nota: str
    celda: Celda


@dataclass
class Modelo:
    fichero: Path
    huella: str
    hojas: List[str]
    titulo: str = ""
    supuestos: List[Supuesto] = field(default_factory=list)
    escenarios: List[Escenario] = field(default_factory=list)
    tabla_escenarios: List[dict] = field(default_factory=list)        # de la hoja resumen: nombre, wacc, g, valor_hoy, valor_rodado, peso
    valor_razonable: Dict[str, float] = field(default_factory=dict)    # «hoy», «rodado»
    referencias_precio: List[dict] = field(default_factory=list)      # rótulo, valor, sobre_entrada, sobre_precio, celda
    multiplos: List[Multiplo] = field(default_factory=list)
    rango_multiplos: Tuple[Optional[float], Optional[float]] = (None, None)
    anclajes_objetivo: List[Tuple[str, float, Celda]] = field(default_factory=list)
    sensibilidad: Optional[dict] = None                                # ejes_wacc, ejes_g, matriz, celda
    reverse: Optional[dict] = None                                     # parametros, anios, filas, resumen
    faltan: Dict[str, str] = field(default_factory=dict)
    recalculado: Optional[Tuple[Path, str]] = None                     # (copia recalculada con Excel, su sha256) si el libro llegó sin valores

    @property
    def nombre(self) -> str:
        return self.fichero.name

    def supuesto(self, patron: str) -> Optional[Supuesto]:
        rx = re.compile(patron, re.I)
        return next((s for s in self.supuestos if rx.search(s.rotulo)), None)

    def escenario(self, patron: str) -> Optional[Escenario]:
        rx = re.compile(patron, re.I)
        return next((e for e in self.escenarios if rx.search(e.nombre)), None)


# ---------------------------------------------------------------------------
# Lectura
# ---------------------------------------------------------------------------

_ANIO = re.compile(r"^(20\d\d)\s*[EPe]?$")
_ESCENARIOS = (("pesimista", r"pesimis|bear|conservador"), ("base", r"^base|central|neutral"), ("optimista", r"optimis|bull"))


def _es_num(v) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _texto(v) -> str:
    return str(v).strip() if isinstance(v, str) else ""


def _sin_valores(valores, formulas) -> int:
    """Cuántas celdas con fórmula no tienen valor calculado guardado (un libro guardado sin recalcular)."""
    n = 0
    for ws in formulas.worksheets:
        wv = valores[ws.title]
        for fila in ws.iter_rows():
            for c in fila:
                if isinstance(c.value, str) and c.value.startswith("=") and wv[c.coordinate].value is None:
                    n += 1
    return n


def recalcular_con_excel(ruta: Path, copia: Path) -> Optional[str]:
    """Abre una COPIA del libro en el Excel del analista (COM, vía PowerShell), recalcula, guarda y devuelve su sha256.

    El libro original no se toca. No es el sistema quien calcula: es el Excel de la máquina evaluando las fórmulas
    del propio analista, lo mismo que ocurriría al abrir el fichero. Sin Excel (u otro sistema), o si tras
    recalcular siguen faltando valores, devuelve None.
    """
    import os
    import shutil
    import subprocess
    if os.name != "nt":
        return None
    copia = Path(copia).resolve()
    copia.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(ruta, copia)
    # el guion va en un fichero .ps1: una ruta con «&» o espacios dentro de -Command no sobrevive al doble entrecomillado de Windows
    guion = copia.with_suffix(".recalcular.ps1")
    lineas = ["$ErrorActionPreference = 'Stop'", "$x = New-Object -ComObject Excel.Application", "$x.Visible = $false", "$x.DisplayAlerts = $false",
              "$wb = $x.Workbooks.Open('" + str(copia).replace("'", "''") + "')", "$x.CalculateFullRebuild()", "$wb.Save()", "$wb.Close($true)", "$x.Quit()",
              "[System.Runtime.InteropServices.Marshal]::ReleaseComObject($x) | Out-Null", "Write-Output listo"]
    guion.write_text(chr(10).join(lineas) + chr(10), encoding="utf-8")
    try:
        r = subprocess.run(["powershell", "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass", "-File", str(guion)],
                           capture_output=True, text=True, timeout=180)
    except (OSError, subprocess.TimeoutExpired):
        return None
    finally:
        guion.unlink(missing_ok=True)
    if r.returncode != 0 or "listo" not in r.stdout:
        return None
    import openpyxl
    if _sin_valores(openpyxl.load_workbook(str(copia), data_only=True), openpyxl.load_workbook(str(copia), data_only=False)):
        return None                      # Excel no dejó valores: no se da por recalculado
    return hashlib.sha256(copia.read_bytes()).hexdigest()


def _contenido(libro) -> dict:
    return {(ws.title, c.coordinate): float(c.value) if isinstance(c.value, (int, float)) and not isinstance(c.value, bool) else c.value
            for ws in libro.worksheets for row in ws.iter_rows() for c in row if c.value is not None}


def _copia_vigente(formulas_original, copia: Path) -> Optional[str]:
    """La copia recalculada de una emisión anterior sirve si existe, está completa y tiene las mismas fórmulas y entradas
    que el original: releer el libro no debe reescribirla, o la huella impresa en el informe dejaría de ser la del fichero."""
    copia = Path(copia)
    if not copia.exists():
        return None
    import openpyxl
    valores, formulas = openpyxl.load_workbook(str(copia), data_only=True), openpyxl.load_workbook(str(copia), data_only=False)
    if _sin_valores(valores, formulas) or _contenido(formulas) != _contenido(formulas_original):
        return None
    return hashlib.sha256(copia.read_bytes()).hexdigest()


class _Libro:
    def __init__(self, ruta: Path, copia_recalculada: Optional[Path] = None):
        import openpyxl
        self.valores = openpyxl.load_workbook(str(ruta), data_only=True)
        self.formulas = openpyxl.load_workbook(str(ruta), data_only=False)
        self.sin_valores = _sin_valores(self.valores, self.formulas)
        self.recalculado: Optional[Tuple[Path, str]] = None
        if self.sin_valores:
            copia = Path(copia_recalculada) if copia_recalculada else Path(tempfile.gettempdir()) / (Path(ruta).stem + ".recalculado.xlsx")
            huella = _copia_vigente(self.formulas, copia) or recalcular_con_excel(Path(ruta), copia)
            if huella is not None:
                self.valores = openpyxl.load_workbook(str(copia), data_only=True)
                self.recalculado = (copia, huella)

    def hoja(self, patron: str):
        rx = re.compile(patron, re.I)
        return next((ws for ws in self.valores.worksheets if rx.search(ws.title)), None)

    def celda(self, ws, ref: str) -> Celda:
        f = self.formulas[ws.title][ref].value
        return Celda(ws.title, ref, f if isinstance(f, str) and f.startswith("=") else None, self.valores[ws.title][ref].number_format)

    @staticmethod
    def filas(ws) -> List[List]:
        return [list(r) for r in ws.iter_rows()]


def _pares(ws, libro: _Libro, max_filas: int = 80) -> List[Tuple[str, float, str, Celda]]:
    """Filas «rótulo · número · texto» de una hoja de parámetros: el rótulo a la izquierda del primer número."""
    salida = []
    for fila in ws.iter_rows(min_row=1, max_row=min(ws.max_row, max_filas)):
        celdas = [c for c in fila if c.value is not None]
        for k, c in enumerate(celdas):
            if _es_num(c.value) and k > 0 and isinstance(celdas[k - 1].value, str):
                origen = " ".join(_texto(x.value) for x in celdas[k + 1:] if isinstance(x.value, str))
                salida.append((celdas[k - 1].value.strip(), float(c.value), origen, libro.celda(ws, c.coordinate)))
                break
    return salida


def _leer_supuestos(libro: _Libro, m: Modelo) -> None:
    ws = libro.hoja(r"supuest|assum|input|par[aá]metr")
    if ws is None:
        m.faltan["supuestos"] = "el libro no tiene hoja de supuestos («Supuestos», «Assumptions», «Inputs»)"
        return
    primera = next((_texto(c.value) for fila in ws.iter_rows(min_row=1, max_row=4) for c in fila if _texto(c.value)), "")
    m.titulo = primera
    for rotulo, valor, origen, celda in _pares(ws, libro):
        m.supuestos.append(Supuesto(rotulo, valor, origen, celda))
    if not m.supuestos:
        m.faltan["supuestos"] = f"la hoja «{ws.title}» no tiene filas rótulo · valor"


def _leer_escenario(libro: _Libro, ws, nombre: str) -> Escenario:
    filas = libro.filas(ws)
    e = Escenario(nombre=nombre, hoja=ws.title, wacc=None, g=None, peso=None, anios=[], filas={}, resumen={})
    # parámetros: la celda numérica a la derecha de cada rótulo
    for fila in filas:
        for k, c in enumerate(fila[:-1]):
            t = _texto(c.value).lower()
            if not t:
                continue
            sig = fila[k + 1].value
            if not _es_num(sig):
                continue
            if t == "wacc" or t.startswith("wacc"):
                e.wacc = float(sig); e.celdas["wacc"] = libro.celda(ws, fila[k + 1].coordinate)
            elif "terminal" in t and ("crecimiento" in t or "growth" in t or t.startswith("g ")):
                e.g = float(sig); e.celdas["g"] = libro.celda(ws, fila[k + 1].coordinate)
            elif (t.startswith("peso") or t.startswith("weight")) and "terminal" not in t:
                e.peso = float(sig); e.celdas["peso"] = libro.celda(ws, fila[k + 1].coordinate)
    # la fila de años y, debajo, las filas con valores en esas columnas
    cabecera = None
    for i, fila in enumerate(filas):
        cols = [(k, _ANIO.match(_texto(c.value)).group(1) + ("E" if _texto(c.value).upper().endswith("E") else "")) for k, c in enumerate(fila) if _texto(c.value) and _ANIO.match(_texto(c.value))]
        if len(cols) >= 3:
            cabecera = (i, cols)
            break
    if cabecera is None:
        e.filas = {}
        return e
    i0, cols = cabecera
    e.anios = [a for _, a in cols]
    col_idx = [k for k, _ in cols]
    for fila in filas[i0 + 1:]:
        rotulo = next((_texto(c.value) for c in fila[:col_idx[0]] if _texto(c.value)), "")
        if not rotulo:
            continue
        valores = [fila[k].value if k < len(fila) else None for k in col_idx]
        numericos = [v for v in valores if _es_num(v)]
        if len(numericos) >= max(2, len(col_idx) - 1):
            e.filas[rotulo] = [float(v) if _es_num(v) else None for v in valores]
            e.celdas[rotulo] = libro.celda(ws, fila[col_idx[0]].coordinate)
            e.citas[rotulo] = [libro.celda(ws, fila[k].coordinate).cita if k < len(fila) else "" for k in col_idx]
        elif len(numericos) == 1 and _es_num(fila[col_idx[0]].value) and rotulo.upper() not in ("WACC",):
            e.resumen[rotulo] = float(fila[col_idx[0]].value)
            e.celdas[rotulo] = libro.celda(ws, fila[col_idx[0]].coordinate)
    return e


def _parece_escenario(v: List[float]) -> bool:
    """Las dos primeras cifras de un escenario son el WACC y el crecimiento terminal: fracciones de 0 a 0,5 (o de 0 a 50
    si el libro los teclea en puntos). Un importe de seis cifras no es un WACC, por mucho que la fila esté bajo el rótulo."""
    wacc, g = v[0], v[1]
    escala = lambda x: abs(x) <= 0.5 or 1 <= abs(x) <= 50
    return escala(wacc) and escala(g) and abs(v[2]) > 0


def _leer_escenarios(libro: _Libro, m: Modelo) -> None:
    for nombre, patron in _ESCENARIOS:
        ws = libro.hoja(patron)
        if ws is None:
            m.faltan[f"escenario_{nombre}"] = f"el libro no tiene hoja del escenario {nombre}"
            continue
        e = _leer_escenario(libro, ws, nombre)
        if not e.filas:
            m.faltan[f"escenario_{nombre}"] = f"la hoja «{ws.title}» no tiene fila de años (2026E…) con proyecciones debajo"
        m.escenarios.append(e)


def _leer_sensibilidad(libro: _Libro, m: Modelo) -> None:
    ws = libro.hoja(r"sensib|sensitiv")
    if ws is None:
        m.faltan["sensibilidad"] = "el libro no tiene hoja de sensibilidad"
        return
    filas = libro.filas(ws)
    for i, fila in enumerate(filas):
        for k, c in enumerate(fila):
            t = _texto(c.value).lower()
            if "wacc" in t and ("g" in t.replace("wacc", "") or "\\" in t or "/" in t):
                ejes_g = [float(x.value) for x in fila[k + 1:] if _es_num(x.value)]
                ejes_wacc, matriz = [], []
                for f2 in filas[i + 1:]:
                    if k < len(f2) and _es_num(f2[k].value):
                        ejes_wacc.append(float(f2[k].value))
                        matriz.append([float(x.value) if _es_num(x.value) else None for x in f2[k + 1:k + 1 + len(ejes_g)]])
                    else:
                        break
                if ejes_g and ejes_wacc:
                    m.sensibilidad = {"ejes_g": ejes_g, "ejes_wacc": ejes_wacc, "matriz": matriz, "celda": libro.celda(ws, c.coordinate)}
                    return
    m.faltan["sensibilidad"] = f"la hoja «{ws.title}» no tiene una matriz con cabecera «WACC \\ g»"


def _leer_reverse(libro: _Libro, m: Modelo) -> None:
    ws = libro.hoja(r"reverse|inverso|impl[ií]cit")
    if ws is None:
        m.faltan["reverse"] = "el libro no tiene hoja de reverse DCF"
        return
    e = _leer_escenario(libro, ws, "reverse")
    parametros = [(r, v, o, c) for r, v, o, c in _pares(ws, libro) if r not in e.resumen and r not in e.filas]
    m.reverse = {"hoja": ws.title, "parametros": parametros, "anios": e.anios, "filas": e.filas, "resumen": e.resumen, "celdas": e.celdas}
    if not e.resumen and not parametros:
        m.faltan["reverse"] = f"la hoja «{ws.title}» no tiene parámetros ni resultados reconocibles"


def _leer_resumen(libro: _Libro, m: Modelo) -> None:
    ws = libro.hoja(r"resumen|summary|valoraci")
    if ws is None:
        m.faltan["resumen"] = "el libro no tiene hoja resumen (escenarios, referencias de precio, múltiplos, precio objetivo)"
        return
    filas = libro.filas(ws)
    bloque = None
    for fila in filas:
        celdas = [c for c in fila if c.value is not None]
        if not celdas:
            continue
        t0 = _texto(celdas[0].value)
        t0l = t0.lower()
        nums = [c for c in celdas[1:] if _es_num(c.value)]
        # una cabecera de bloque es una fila sin números; «Precio objetivo del inversor 100» es un dato, no una cabecera
        if not nums:
            if t0l.startswith("escenario") or t0l.startswith("scenario"):
                bloque = "escenarios"
            elif "referencias de precio" in t0l or "price references" in t0l:
                bloque = "precio"
            elif "múltiplos" in t0l or "multiplos" in t0l or "multiples" in t0l:
                bloque = "multiplos"
            elif t0l.startswith("precio objetivo") or t0l.startswith("target price"):
                bloque = "objetivo"
            continue
        if bloque == "escenarios":
            if "razonable" in t0l or "fair value" in t0l:
                vals = [float(c.value) for c in nums]
                if vals:
                    m.valor_razonable = {"hoy": vals[0], "rodado": vals[1] if len(vals) > 1 else None, "celda": libro.celda(ws, nums[0].coordinate)}
                bloque = None
            elif len(nums) >= 4:
                v = [float(c.value) for c in nums]
                if not _parece_escenario(v):
                    # la fila está bajo una cabecera que dice «escenario» pero no es un escenario (un «Equity Value»,
                    # una fila de acciones): imprimirla daría un WACC de siete cifras. Se anota y no se toma.
                    m.faltan.setdefault("tabla_escenarios", f"la hoja «{ws.title}» tiene un bloque de escenarios cuyas filas no traen "
                                                            f"WACC y crecimiento reconocibles (p. ej. «{t0}»): no se lee")
                    continue
                m.tabla_escenarios.append({"nombre": t0, "wacc": v[0], "g": v[1], "valor_hoy": v[2], "valor_rodado": v[3] if len(v) > 4 else None,
                                           "peso": v[4] if len(v) > 4 else v[3], "celda": libro.celda(ws, nums[2].coordinate),
                                           "celdas": [libro.celda(ws, c.coordinate) for c in nums]})
        elif bloque == "precio" and nums:
            v = [float(c.value) for c in nums]
            m.referencias_precio.append({"rotulo": t0, "valor": v[0], "sobre_entrada": v[1] if len(v) > 1 else None,
                                         "sobre_precio": v[2] if len(v) > 2 else None, "celda": libro.celda(ws, nums[0].coordinate),
                                         "celdas": [libro.celda(ws, c.coordinate) for c in nums]})
        elif bloque == "multiplos":
            if "mínimo" in t0l or "minimo" in t0l or "min" == t0l[:3]:
                m.rango_multiplos = (float(nums[0].value) if nums else None, m.rango_multiplos[1]); continue
            if "máximo" in t0l or "maximo" in t0l or "max" == t0l[:3]:
                m.rango_multiplos = (m.rango_multiplos[0], float(nums[0].value) if nums else None); continue
            if len(nums) >= 3:
                v = [float(c.value) for c in nums]
                textos = [_texto(c.value) for c in celdas[1:] if isinstance(c.value, str)]
                m.multiplos.append(Multiplo(rotulo=t0, base=v[0], actual=v[1], objetivo=v[2], valor_accion=v[3] if len(v) > 3 else None,
                                            origen=" ".join(textos), celda=libro.celda(ws, nums[-1].coordinate),
                                            celdas=[libro.celda(ws, c.coordinate) for c in nums]))
        elif bloque == "objetivo" and nums:
            m.anclajes_objetivo.append((t0, float(nums[0].value), libro.celda(ws, nums[0].coordinate)))
    if not m.tabla_escenarios:
        m.faltan["tabla_escenarios"] = f"la hoja «{ws.title}» no tiene la tabla «Escenario · WACC · g · valor · peso»"
    if not m.multiplos:
        m.faltan["multiplos"] = f"la hoja «{ws.title}» no tiene el bloque «Valoración por múltiplos»"
    if not m.anclajes_objetivo:
        m.faltan["precio_objetivo"] = f"la hoja «{ws.title}» no tiene el bloque «Precio objetivo»"


def cargar(ruta: Path, copia_recalculada: Optional[Path] = None) -> Modelo:
    """Lee el libro. Si se guardó sin recalcular (fórmulas sin valor), recalcula una copia con el Excel de la máquina;
    sin Excel, lo que son fórmulas sale N/A y `faltan` lo dice."""
    ruta = Path(ruta)
    libro = _Libro(ruta, copia_recalculada)
    m = Modelo(fichero=ruta, huella=hashlib.sha256(ruta.read_bytes()).hexdigest(), hojas=libro.valores.sheetnames)
    if libro.sin_valores:
        if libro.recalculado is not None:
            m.recalculado = libro.recalculado
            m.faltan["valores"] = (f"el libro se guardó sin recalcular ({libro.sin_valores} celdas de fórmula sin valor); los valores impresos son los de una copia "
                                   f"recalculada con el Excel de la máquina ({m.recalculado[0].name}, sha256 {m.recalculado[1][:12]}); el original no se ha tocado")
        else:
            m.faltan["valores"] = (f"el libro se guardó sin recalcular ({libro.sin_valores} celdas de fórmula sin valor) y no hay Excel con que recalcular una copia: "
                                   "las cifras que son fórmula salen N/A; ábrelo en Excel, recalcula y guarda")
    _leer_supuestos(libro, m)
    _leer_escenarios(libro, m)
    _leer_sensibilidad(libro, m)
    _leer_reverse(libro, m)
    _leer_resumen(libro, m)
    return m


# ---------------------------------------------------------------------------
# Cuadre con los hechos contrastados (regla 9 aplicada al modelo del analista)
# ---------------------------------------------------------------------------

_FECHA = re.compile(r"(\d{1,2})/(\d{1,2})/(20\d\d)")
# rótulo del supuesto → (campo del informe, tipo, escala del libro respecto al hecho)
_CUADRES = (
    (re.compile(r"^ingresos.*?(20\d\d)", re.I), "ingresos", "anual", 1e6),
    (re.compile(r"^margen operativo.*?(20\d\d)", re.I), "margen_ebit", "anual", 1.0),
    (re.compile(r"deuda financiera|deuda total|deuda bruta|gross debt", re.I), "deuda_bruta", "instante", 1e6),
    (re.compile(r"^caja", re.I), "caja_total", "instante", 1e6),
    (re.compile(r"acciones en circulaci", re.I), "acciones_circulacion", "instante", 1e6),
    (re.compile(r"acciones diluidas", re.I), "acciones_diluidas", "trimestre", 1e6),
)


def _recorte(texto: str, maximo: int = 200) -> str:
    """Un texto del analista, entero si cabe y, si no, cortado en la última frase o palabra completa antes del límite."""
    texto = " ".join(texto.split())
    if len(texto) <= maximo:
        return texto
    corte = texto.rfind(". ", 0, maximo)
    if corte < maximo // 2:
        corte = texto.rfind(" ", 0, maximo)
    return texto[:corte].rstrip(".") + "…"


def _fecha_de(texto: str) -> Optional[date]:
    m = _FECHA.search(texto)
    return date(int(m.group(3)), int(m.group(2)), int(m.group(1))) if m else None


def fecha_precio_libro(m: Modelo) -> Optional[date]:
    """La fecha del precio de mercado que el analista tecleó, si la escribió junto a él («Cierre del 14/09/2026…»)."""
    s = m.supuesto("precio de mercado")
    return (_fecha_de(s.origen) or _fecha_de(s.rotulo)) if s is not None else None


def cuadrar(m: Modelo, hechos: Dict[Tuple[str, Periodo], Hecho], fin_ejercicio: Optional[date] = None,
            cierres: Optional[Dict[date, float]] = None, agregador=None) -> List[Cuadre]:
    """Cada supuesto del libro que afirma ser un hecho, frente al hecho contrastado del informe; y el precio
    de mercado que el analista tecleó, frente al cierre oficial de esa fecha (`cierres`, del histórico de la bolsa).

    Cuando difieren, el informe se rige por el dato oficial (lo dice el cuadro) y la nota explica la diferencia
    con lo que las fuentes permiten afirmar: si la cifra del libro coincide con la «deuda total» del agregador
    (`agregador`), que incluye pasivos por arrendamiento, se dice; si el libro declara una cifra propia, se cita.
    """
    salida: List[Cuadre] = []
    s = m.supuesto("precio de mercado")
    if s is not None:
        f = fecha_precio_libro(m)
        if f is None:
            salida.append(Cuadre(s.rotulo, s.valor, "precio", "?", None, Contraste.HUECO, "el libro no dice de qué fecha es su precio de mercado", s.celda))
        elif not cierres:
            salida.append(Cuadre(s.rotulo, s.valor, "precio", f"{f:%d/%m/%Y}", None, Contraste.HUECO,
                                 "sin histórico oficial de cierres con que cuadrarlo (fuente de cotización sin configurar o sin respuesta)", s.celda))
        elif f not in cierres:
            salida.append(Cuadre(s.rotulo, s.valor, "precio", f"{f:%d/%m/%Y}", None, Contraste.HUECO,
                                 f"el histórico oficial no tiene cierre para el {f:%d/%m/%Y} (¿no fue día de mercado?)", s.celda))
        else:
            oficial = cierres[f]
            if abs(oficial - s.valor) <= max(abs(oficial) * 0.005, 0.005):
                salida.append(Cuadre(s.rotulo, s.valor, "precio", f"{f:%d/%m/%Y}", oficial, Contraste.CONFIRMADO,
                                     f"coincide con el cierre oficial del {f:%d/%m/%Y} ({numero(oficial, 2)} USD, histórico de la bolsa)", s.celda))
            else:
                salida.append(Cuadre(s.rotulo, s.valor, "precio", f"{f:%d/%m/%Y}", oficial, Contraste.DISCREPANTE,
                                     f"el libro dice {numero(s.valor, 2)} y el cierre oficial del {f:%d/%m/%Y} fue {numero(oficial, 2)} USD; rige el cierre oficial (el libro cita: {_recorte(s.origen)})", s.celda))
    instantes = sorted({p for (_, p) in hechos if p.es_instante})
    trimestres = sorted({p for (_, p) in hechos if not p.es_instante and p.meses == 3})
    for s in m.supuestos:
        for rx, campo, tipo, escala in _CUADRES:
            mm = rx.search(s.rotulo)
            if not mm:
                continue
            periodo: Optional[Periodo] = None
            if tipo == "anual":
                anio = int(mm.group(1))
                periodo = next((p for (c, p) in hechos if c == "ingresos" and not p.es_instante and p.meses == 12 and p.fin.year == anio), None)
            elif tipo == "instante":
                f = _fecha_de(s.origen) or _fecha_de(s.rotulo)
                periodo = Periodo.instante(f) if f else (instantes[-1] if instantes else None)
            else:
                periodo = trimestres[-1] if trimestres else None
            if periodo is None:
                salida.append(Cuadre(s.rotulo, s.valor, campo, "?", None, Contraste.HUECO, "el informe no tiene periodo con que cuadrar este supuesto", s.celda))
                break
            if campo == "caja_total":
                caja, inv = hechos.get(("caja", periodo)), hechos.get(("inversiones_cp", periodo))
                valor = (caja.valor if caja and caja.hay_dato else None)
                if valor is not None and inv is not None and inv.hay_dato:
                    valor += inv.valor
                rotulo_informe = "caja + inversiones a corto"
            else:
                h = hechos.get((campo, periodo))
                valor = h.valor if h is not None and h.hay_dato else None
                rotulo_informe = campo
            if valor is None:
                salida.append(Cuadre(s.rotulo, s.valor, campo, periodo.clave, None, Contraste.HUECO,
                                     f"el informe no tiene {rotulo_informe} para {periodo.clave}", s.celda))
                break
            en_libro = valor / escala
            tol = max(abs(en_libro) * 0.005, 0.5 if escala > 1 else 0.0005)
            if abs(en_libro - s.valor) <= tol:
                salida.append(Cuadre(s.rotulo, s.valor, campo, periodo.clave, en_libro, Contraste.CONFIRMADO,
                                     f"coincide con {rotulo_informe} {periodo.clave} (sección C)", s.celda))
            else:
                dif = (s.valor - en_libro) / en_libro if en_libro else None
                fmt = (lambda v: numero(v, 0)) if abs(en_libro) >= 100 else (lambda v: numero(v, 4))
                explicacion = ""
                deuda_agregador = getattr(agregador, "deuda_total", None) if agregador is not None else None
                if campo == "deuda_bruta" and deuda_agregador and abs(deuda_agregador / escala - s.valor) <= max(abs(s.valor) * 0.005, 0.5):
                    explicacion = (f"; la cifra del libro coincide con la «deuda total» que publica el agregador ({fmt(deuda_agregador / escala)}), que incluye "
                                   "pasivos por arrendamiento; la deuda financiera de la SEC (corto + largo plazo) es la que rige")
                elif s.origen:
                    explicacion = f"; el libro la justifica así: «{_recorte(s.origen)}»"
                salida.append(Cuadre(s.rotulo, s.valor, campo, periodo.clave, en_libro, Contraste.DISCREPANTE,
                                     f"el libro dice {fmt(s.valor)} y el dato oficial más reciente ({periodo.clave}, SEC) es {fmt(en_libro)}{f' ({numero(dif * 100, 1)} %)' if dif is not None else ''}; rige el oficial{explicacion}", s.celda))
            break
    return salida
