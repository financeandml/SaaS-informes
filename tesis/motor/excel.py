"""Excel (05 §9): exportar el motor con fórmulas vivas e importar el libro del analista solo para comparar.

Exportar: hojas Supuestos, Resumen, Pesimista, Base, Optimista, Sensibilidad y Reverse DCF; entradas en amarillo;
columna «Origen / justificación»; rangos con nombre. Recalculado (Excel o LibreOffice sin interfaz), da lo mismo que
el motor ± 0,01.

Importar: solo por rangos con nombre o, en libros antiguos, por un mapa explícito de celdas
(`config/excel_mapas/<nombre>.yaml`). Nunca por rótulo. Rige el motor; el libro solo aparece en la comparación.
"""

from __future__ import annotations

import hashlib
import shutil
import subprocess
import tempfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, List, Optional

from .escenarios import Valoracion
from .supuestos import NOMBRES
from ..rutas import CONFIG

__all__ = ["exportar", "recalcular", "importar", "Comparacion", "comparar", "CORTOS"]

CORTOS = {"pesimista": "pes", "base": "base", "optimista": "opt"}
_AMARILLO = "FFF2CC"


def _col(i: int) -> str:
    from openpyxl.utils import get_column_letter
    return get_column_letter(i)


def exportar(v: Valoracion, ingresos_base: float, ruta: Path, umbrales: dict) -> Path:
    """Escribe el libro con fórmulas vivas que reproducen el motor."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.workbook.defined_name import DefinedName
    wb = Workbook()
    amarillo = PatternFill("solid", fgColor=_AMARILLO)
    negrita = Font(bold=True)

    def nombre(clave: str, hoja: str, celda: str) -> None:
        ref = f"'{hoja}'!${''.join(c for c in celda if c.isalpha())}${''.join(c for c in celda if c.isdigit())}"
        wb.defined_names[clave] = DefinedName(clave, attr_text=ref)

    su = wb.active
    su.title = "Supuestos"
    p, pte, w = v.parametros, v.puente, v.wacc
    base = v.base
    filas = [
        ("precio", "Precio: cierre oficial en la fecha de valoración (USD/acción)", v.precio, "Nasdaq"),
        ("acciones_diluidas", "Acciones diluidas", pte.acciones, pte.acciones_nota),
        ("ingresos_base", "Ingresos del año base (USD)", ingresos_base, "SEC (hechos verificados)"),
        ("fraccion", "Fracción del ejercicio 1 que queda (f)", base.proyeccion.fraccion, "días hasta el cierre / días del ejercicio"),
        ("fraccion_flujo", "Parte del FCFF del año 1 que entra (f_flujo)", base.proyeccion.fraccion_flujo,
         "días del último balance al cierre / días del ejercicio"),
        ("mitad_de_anio", "Convención de mitad de año (1 sí, 0 no)", 1 if p.mitad_de_anio else 0, "entradas del analista"),
        ("bin_inicial", "Bases imponibles negativas iniciales (USD)", p.bin_inicial, "entradas del analista"),
        ("ke", "Coste de los fondos propios (Ke)", w.ke, "rf + β × ERP + prima"),
        ("horizonte_meses", "Horizonte (meses)", p.horizonte_meses, "entradas del analista"),
        ("dpa_horizonte", "Dividendos por acción esperados en el horizonte (USD)", v.dpa_horizonte, "DPA anual vigente × horizonte / 12: bolsa, SEC o el cero que declara el 10-K (fuente en el informe)"),
        ("ajuste_puente", "Ajuste del puente: − deuda + caja + inversiones ± ajustes (USD)", pte.ajuste, "último balance publicado"),
        ("deuda_neta", "Deuda neta (USD)", pte.deuda_neta, "último balance publicado"),
        ("wacc_calculado", "WACC calculado (antes del ajuste por escenario)", w.wacc, "E/(D+E) × Ke + D/(D+E) × Kd × (1 − t)"),
        ("sbc_coste", "SBC como coste de caja (1 sí, 0 no)", 1 if p.sbc_politica == "coste_de_caja" else 0, "entradas del analista"),
    ]
    su["A1"], su["B1"], su["C1"] = "Rango", "Supuesto", "Valor"
    su["D1"] = "Origen / justificación"
    for i, (clave, rotulo, valor, origen) in enumerate(filas, 2):
        su[f"A{i}"], su[f"B{i}"], su[f"C{i}"], su[f"D{i}"] = clave, rotulo, valor, origen
        su[f"C{i}"].fill = amarillo
        nombre(clave, "Supuestos", f"C{i}")
    hojas = {"pesimista": "Pesimista", "base": "Base", "optimista": "Optimista"}
    n = len(base.escenario.crecimiento)
    ultima = _col(2 + n)
    for clave_esc, titulo in hojas.items():
        r = v.resultados.get(clave_esc)
        if r is None:
            continue
        e = r.escenario
        ws = wb.create_sheet(titulo)
        c = CORTOS[clave_esc]
        cab = [("wacc", "WACC del escenario", r.wacc), ("g", "Crecimiento terminal (g)", e.g), ("peso", "Probabilidad", e.probabilidad),
               ("ronic", "RONIC", e.ronic if e.ronic is not None else None), ("multiplo", "Múltiplo de salida", e.multiplo_salida),
               ("fm", "ΔFM / Δingresos", e.fm)]
        for i, (k, rotulo, valor) in enumerate(cab, 2):
            ws[f"A{i}"], ws[f"B{i}"] = f"{k}_{c}", rotulo
            if k == "ronic" and valor is None:
                ws[f"C{i}"] = f"=C2+{umbrales['ronic_defecto_pp'] / 100}"
            else:
                ws[f"C{i}"] = valor
                ws[f"C{i}"].fill = amarillo
            nombre(f"{k}_{c}", titulo, f"C{i}")
        ws["B9"] = "Año"
        for t in range(n):
            ws.cell(9, 3 + t, f"Año {t + 1}")
        drivers = [("crecimiento", e.crecimiento), ("margen", e.margen), ("impuesto", e.impuesto), ("da", e.da), ("capex", e.capex),
                   ("sbc", e.sbc)] + [(k, s) for k, s in e.paquete.items()]
        for j, (rot, serie) in enumerate(drivers):
            fila = 10 + j
            ws.cell(fila, 2, rot)
            for t, x in enumerate(serie):
                ws.cell(fila, 3 + t, x).fill = amarillo
        base_fila = 10 + len(drivers) + 1
        F = {k: base_fila + i for i, k in enumerate(("ingresos", "ebit", "bin", "impuestos", "nopat", "da", "capex", "dfm", "sbc",
                                                        "paquete", "fcff", "ebitda", "t", "df", "peso_anio", "va"))}
        rotulos = {"ingresos": "Ingresos", "ebit": "EBIT", "bin": "Bases imponibles negativas al inicio", "impuestos": "Impuestos en caja",
                   "nopat": "NOPAT", "da": "D&A", "capex": "Capex", "dfm": "ΔFM", "sbc": "SBC", "paquete": "Partidas del paquete",
                   "fcff": "FCFF", "ebitda": "EBITDA", "t": "Momento de descuento (años)", "df": "Factor de descuento",
                   "peso_anio": "Parte del año que cuenta", "va": "Valor actual del FCFF"}
        for k, fila in F.items():
            ws.cell(fila, 2, rotulos[k])
        paquete_filas = [10 + 6 + j for j in range(len(e.paquete))]
        for t in range(n):
            col = _col(3 + t)
            prev = _col(2 + t)
            ref = lambda k: f"{col}{F[k]}"                                         # noqa: E731
            ws[ref("ingresos")] = f"=ingresos_base*(1+{col}10)" if t == 0 else f"={prev}{F['ingresos']}*(1+{col}10)"
            ws[ref("ebit")] = f"={ref('ingresos')}*{col}11+(1-sbc_coste)*{ref('sbc')}"   # «dilución»: SBC sumada al EBIT
            ws[ref("bin")] = "=bin_inicial" if t == 0 else f"=MAX(0,{prev}{F['bin']}-MAX({prev}{F['ebit']},0))+MAX(-{prev}{F['ebit']},0)"
            ws[ref("impuestos")] = f"=MAX(0,{ref('ebit')}-{ref('bin')})*{col}12"
            ws[ref("nopat")] = f"={ref('ebit')}-{ref('impuestos')}"
            ws[ref("da")] = f"={ref('ingresos')}*{col}13"
            ws[ref("capex")] = f"={ref('ingresos')}*{col}14"
            ws[ref("dfm")] = (f"=({ref('ingresos')}-ingresos_base)*fm_{c}" if t == 0
                             else f"=({ref('ingresos')}-{prev}{F['ingresos']})*fm_{c}")
            ws[ref("sbc")] = f"={ref('ingresos')}*{col}15"                            # informativa con «coste de caja»
            ws[ref("paquete")] = ("=" + "+".join(f"{ref('ingresos')}*{col}{f}" for f in paquete_filas)) if paquete_filas else 0
            ws[ref("fcff")] = f"={ref('nopat')}+{ref('da')}-{ref('capex')}-{ref('dfm')}-{ref('paquete')}"
            ws[ref("ebitda")] = f"={ref('ebit')}+{ref('da')}"
            ws[ref("t")] = ("=IF(mitad_de_anio=1,fraccion-fraccion_flujo/2,fraccion)" if t == 0 else f"=fraccion+{t}-IF(mitad_de_anio=1,0.5,0)")
            ws[ref("df")] = f"=(1+wacc_{c})^(-{ref('t')})"
            ws[ref("peso_anio")] = "=fraccion_flujo" if t == 0 else 1
            ws[ref("va")] = f"={ref('fcff')}*{ref('df')}*{ref('peso_anio')}"
        fin = F["va"] + 2
        ult = f"{ultima}"
        vt = {"value_driver": f"={ult}{F['nopat']}*(1+g_{c})*(1-g_{c}/ronic_{c})/(wacc_{c}-g_{c})",
              "gordon": f"={ult}{F['fcff']}*(1+g_{c})/(wacc_{c}-g_{c})",
              "multiplo_salida": f"={ult}{F['ebitda']}*multiplo_{c}"}[r.terminal.metodo]
        resumen = [("suma_va", "Suma del valor actual de los FCFF", f"=SUM(C{F['va']}:{ult}{F['va']})"),
                   ("vt", f"Valor terminal ({r.terminal.metodo})", vt),
                   ("va_vt", "Valor actual del valor terminal", f"=C{fin + 1}*(1+wacc_{c})^(-(fraccion+{n}-1))"),
                   ("ev", "Valor de empresa", f"=C{fin}+C{fin + 2}"),
                   ("fondos_propios", "Fondos propios", f"=C{fin + 3}+ajuste_puente"),
                   ("acciones_valor", "Acciones del valor por acción (+ las que pagan la SBC con «dilución»)",
                    f"=acciones_diluidas+(1-sbc_coste)*SUMPRODUCT(C{F['sbc']}:{ult}{F['sbc']},C{F['peso_anio']}:{ult}{F['peso_anio']})/precio"),
                   ("valor_accion", "Valor por acción hoy (V₀)", f"=C{fin + 4}/C{fin + 5}"),
                   ("valor_h", "Valor por acción en el horizonte (V_h)", f"=C{fin + 6}*(1+ke)^(horizonte_meses/12)-dpa_horizonte")]
        for i, (k, rot, formula) in enumerate(resumen):
            ws[f"A{fin + i}"], ws[f"B{fin + i}"], ws[f"C{fin + i}"] = f"{k}_{c}", rot, formula
            ws[f"B{fin + i}"].font = negrita
            nombre(f"{k}_{c}", titulo, f"C{fin + i}")
        ws["A1"] = f"Escenario {clave_esc}: fórmulas vivas del motor (05 §3–§7)"
        # para la sensibilidad: filas que se reutilizan
        if clave_esc == "base":
            base_filas = F
    re = wb.create_sheet("Resumen", 1)
    re["A1"], re["B1"], re["C1"], re["D1"], re["E1"] = "Escenario", "Probabilidad", "V₀", "V_h", "WACC"
    for i, clave_esc in enumerate(NOMBRES, 2):
        c = CORTOS[clave_esc]
        if clave_esc not in v.resultados:
            continue
        re[f"A{i}"], re[f"B{i}"], re[f"C{i}"], re[f"D{i}"], re[f"E{i}"] = clave_esc, f"=peso_{c}", f"=valor_accion_{c}", f"=valor_h_{c}", f"=wacc_{c}"
    metricas = [("valor_razonable", "Valor razonable hoy (Σ p·V₀)", "=SUMPRODUCT(B2:B4,C2:C4)"),
                ("po", "Precio objetivo (Σ p·V_h)", "=SUMPRODUCT(B2:B4,D2:D4)"),
                ("potencial", "Potencial", "=po/precio-1"),
                ("margen_seguridad", "Margen de seguridad", "=1-precio/valor_razonable"),
                ("downside", "Downside pesimista", "=valor_h_pes/precio-1"),
                ("recorrido_riesgo", "Recorrido / riesgo", '=IF(valor_h_pes>=precio,"sin pérdida",(po-precio)/(precio-valor_h_pes))')]
    for i, (k, rot, formula) in enumerate(metricas, 7):
        re[f"A{i}"], re[f"B{i}"], re[f"C{i}"] = k, rot, formula
        nombre(k, "Resumen", f"C{i}")
    # sensibilidad WACC × g del base: los FCFF no dependen del WACC; se descuentan con el de cada celda
    se = wb.create_sheet("Sensibilidad")
    b = v.base
    pasos_w, pasos_g = umbrales["sensibilidad"]["wacc_pp"], umbrales["sensibilidad"]["g_pp"]
    se["A1"] = "V₀ del escenario base: WACC (filas) × g (columnas)"
    F = base_filas
    rng = lambda k: f"Base!$C${F[k]}:${ultima}${F[k]}"                            # noqa: E731
    for j, dg in enumerate(pasos_g):
        se.cell(2, 2 + j, b.escenario.g + dg / 100)
    for i, dw in enumerate(pasos_w):
        se.cell(3 + i, 1, b.wacc + dw / 100)
        for j in range(len(pasos_g)):
            wc, gc = f"$A{3 + i}", f"{_col(2 + j)}$2"
            ronic = f"{b.escenario.ronic}" if b.escenario.ronic is not None else f"({wc}+{umbrales['ronic_defecto_pp'] / 100})"
            vt = {"value_driver": f"Base!${ultima}${F['nopat']}*(1+{gc})*(1-{gc}/{ronic})/({wc}-{gc})",
                  "gordon": f"Base!${ultima}${F['fcff']}*(1+{gc})/({wc}-{gc})",
                  "multiplo_salida": f"Base!${ultima}${F['ebitda']}*multiplo_base"}[b.terminal.metodo]
            se.cell(3 + i, 2 + j, f"=IF({wc}-{gc}<=0,\"\",(SUMPRODUCT({rng('fcff')},{rng('peso_anio')},(1+{wc})^(-{rng('t')}))"
                                  f"+{vt}*(1+{wc})^(-(fraccion+{n}-1))+ajuste_puente)/acciones_valor_base)")
    rv = wb.create_sheet("Reverse DCF")
    rv["A1"] = "DCF inverso (motor): qué hay que creer para que V₀ del base sea el precio"
    wb.save(ruta)
    return Path(ruta)


def _sin_valores(valores, formulas) -> int:
    """Cuántas celdas con fórmula no tienen valor calculado guardado (un libro guardado sin recalcular)."""
    n = 0
    for ws in formulas.worksheets:
        wv = valores[ws.title]
        for fila in ws.iter_rows():
            for c in fila:
                if isinstance(c.value, str) and c.value.startswith("=") and wv[c.coordinate].value is None:
                    n += 1
    return n


def recalcular_con_excel(ruta: Path, copia: Path) -> Optional[str]:
    """Abre una COPIA del libro en el Excel del analista (COM, vía PowerShell), recalcula, guarda y devuelve su sha256.

    El libro original no se toca. No es el sistema quien calcula: es el Excel de la máquina evaluando las fórmulas
    del propio analista, lo mismo que ocurriría al abrir el fichero. Sin Excel (u otro sistema), o si tras
    recalcular siguen faltando valores, devuelve None.
    """
    import os
    import shutil
    import subprocess
    if os.name != "nt":
        return None
    copia = Path(copia).resolve()
    copia.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(ruta, copia)
    # el guion va en un fichero .ps1: una ruta con «&» o espacios dentro de -Command no sobrevive al doble entrecomillado de Windows
    guion = copia.with_suffix(".recalcular.ps1")
    lineas = ["$ErrorActionPreference = 'Stop'", "$x = New-Object -ComObject Excel.Application", "$x.Visible = $false", "$x.DisplayAlerts = $false",
              "$wb = $x.Workbooks.Open('" + str(copia).replace("'", "''") + "')", "$x.CalculateFullRebuild()", "$wb.Save()", "$wb.Close($true)", "$x.Quit()",
              "[System.Runtime.InteropServices.Marshal]::ReleaseComObject($x) | Out-Null", "Write-Output listo"]
    guion.write_text(chr(10).join(lineas) + chr(10), encoding="utf-8")
    try:
        r = subprocess.run(["powershell", "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass", "-File", str(guion)],
                           capture_output=True, text=True, timeout=180)
    except (OSError, subprocess.TimeoutExpired):
        return None
    finally:
        guion.unlink(missing_ok=True)
    if r.returncode != 0 or "listo" not in r.stdout:
        return None
    import openpyxl
    if _sin_valores(openpyxl.load_workbook(str(copia), data_only=True), openpyxl.load_workbook(str(copia), data_only=False)):
        return None                      # Excel no dejó valores: no se da por recalculado
    return hashlib.sha256(copia.read_bytes()).hexdigest()


def recalcular(ruta: Path, destino: Path) -> Optional[Path]:
    """Una copia recalculada: Excel en Windows (el del analista) o LibreOffice sin interfaz. Sin ninguno, None."""
    destino = Path(destino)
    if recalcular_con_excel(Path(ruta), destino):
        return destino
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return None
    with tempfile.TemporaryDirectory() as tmp:
        r = subprocess.run([soffice, "--headless", "--calc", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", tmp, str(ruta)],
                           capture_output=True, text=True, timeout=180)
        salida = Path(tmp) / Path(ruta).name
        if r.returncode != 0 or not salida.exists():
            return None
        destino.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(salida, destino)
    return destino


def _mapa(nombre: str) -> Dict[str, str]:
    import yaml
    ruta = CONFIG / "excel_mapas" / f"{nombre}.yaml"
    if not ruta.exists():
        return {}
    return dict((yaml.safe_load(ruta.read_text(encoding="utf-8")) or {}).get("celdas") or {})


def importar(ruta: Path, mapa: Optional[str] = None, recalculado: Optional[Path] = None) -> Dict[str, object]:
    """{rango: valor} leídos por nombre (rangos definidos en el libro) o por el mapa explícito. Nunca por rótulo.
    Con `recalculado`, los valores salen de esa copia (el original puede no guardar valores)."""
    import openpyxl
    formulas = openpyxl.load_workbook(str(ruta), data_only=False)
    valores = openpyxl.load_workbook(str(recalculado or ruta), data_only=True)
    salida: Dict[str, object] = {}
    for clave, dn in formulas.defined_names.items():
        for hoja, celda in dn.destinations:
            salida[clave] = valores[hoja][celda.replace("$", "")].value
    if not salida and mapa:
        for clave, ref in _mapa(mapa).items():
            hoja, celda = ref.rsplit("!", 1)
            salida[clave] = valores[hoja.strip("'")][celda].value
    return salida


@dataclass
class Comparacion:
    filas: List[tuple] = field(default_factory=list)      # (concepto, libro, motor, unidad)
    faltan: List[str] = field(default_factory=list)


def comparar(libro: Dict[str, object], v: Valoracion) -> Comparacion:
    motor = {"precio": v.precio, "acciones_diluidas": v.puente.acciones / 1e6, "deuda_neta": v.puente.deuda_neta / 1e6,
             "valor_razonable": v.valor_razonable, "po": v.po}
    for nombre, c in CORTOS.items():
        r = v.resultados.get(nombre)
        if r is not None:
            motor.update({f"wacc_{c}": r.wacc, f"g_{c}": r.escenario.g, f"peso_{c}": r.escenario.probabilidad, f"valor_accion_{c}": r.v0})
    unidades = {"precio": "USD", "acciones_diluidas": "M", "deuda_neta": "M USD", "valor_razonable": "USD", "po": "USD"}
    rotulos = {"precio": "Precio", "acciones_diluidas": "Acciones diluidas", "deuda_neta": "Deuda neta", "valor_razonable": "Valor razonable hoy",
               "po": "Precio objetivo", "objetivo": "Precio objetivo del libro"}
    comp = Comparacion()
    for clave in ("precio", "acciones_diluidas", "deuda_neta", "wacc_pes", "wacc_base", "wacc_opt", "g_pes", "g_base", "g_opt",
                  "peso_pes", "peso_base", "peso_opt", "valor_accion_pes", "valor_accion_base", "valor_accion_opt", "valor_razonable", "po"):
        if clave not in motor:
            continue
        valor = libro.get(clave)
        if not isinstance(valor, (int, float)) or isinstance(valor, bool):
            comp.faltan.append(f"el libro no da «{clave}» (ni rango con nombre ni celda en el mapa)")
            valor = None
        rot = rotulos.get(clave) or {"wacc": "WACC", "g": "g terminal", "peso": "Probabilidad", "valor": "V₀"}[clave.split("_")[0]] + \
            f" ({clave.rsplit('_', 1)[-1]})"
        unidad = unidades.get(clave, "%" if clave.startswith(("wacc", "g_", "peso")) else "USD")
        comp.filas.append((rot, valor, motor[clave], unidad))
    return comp
